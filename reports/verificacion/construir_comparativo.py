"""Build comparison strictly from recorded original-Excel and SIDE observations."""
import json
from copy import copy
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'reports' / 'verificacion'
THEMES = ['DOP por cantidad', 'costo de producción', 'caja', 'consecuencias', 'puntajes']
NAVY='142E3F'; TEAL='087E8B'; PALE='E9F4F5'; GRAY='647580'

def read(name):
    return json.loads((OUT/name).read_text(encoding='utf-8-sig'))

def theme(value):
    v=str(value).lower()
    if 'dop' in v or 'cantidad' in v: return THEMES[0]
    if 'costo' in v or 'producci' in v: return THEMES[1]
    if 'caja' in v: return THEMES[2]
    if 'consec' in v: return THEMES[3]
    if 'punt' in v: return THEMES[4]
    raise ValueError(f'Tema desconocido: {value}')

def short(value):
    return json.dumps(value, ensure_ascii=False, separators=(', ', ': ')) if isinstance(value,(dict,list)) else str(value)

def style(ws, widths):
    ws.sheet_view.showGridLines=False
    for row in ws:
        for c in row:
            c.font=Font(name='Arial',size=10,color='172D3A')
            c.alignment=Alignment(vertical='top',wrap_text=True)
    for col,width in widths.items(): ws.column_dimensions[col].width=width
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.print_options.horizontalCentered=True
    ws.oddFooter.center.text='SIDE / EXCEL · Página &P de &N'
    ws.oddFooter.center.size=9
    ws.freeze_panes='E8'

def title(ws, text, end='J'):
    ws.merge_cells(f'A1:{end}2'); ws['A1']=text
    ws['A1'].font=Font(name='Arial',size=20,bold=True,color='FFFFFF')
    ws['A1'].fill=PatternFill('solid',fgColor=NAVY)
    ws['A1'].alignment=Alignment(vertical='center')
    ws.row_dimensions[1].height=24; ws.row_dimensions[2].height=15

def side_index(data):
    result={}
    for scenario in data.get('escenarios',[]):
        for m in scenario.get('mediciones',[]): result[(scenario['id'],m['key'])]=m
    return result

def build():
    source=read('datos_excel.json'); observed=read('datos_side.json')
    index=side_index(observed)
    wb=Workbook(); summary=wb.active; summary.title='Resumen'
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
    sheets={name:wb.create_sheet(name) for name in THEMES}
    provenance=wb.create_sheet('Fuentes originales')
    title(summary,'Verificación manual · SIDE vs Excel','H')
    summary['A4']='Tolerancia absoluta'; summary['B4']=source.get('tolerancia_absoluta',0.01)
    summary['B4'].number_format='0.000000'; summary['B4'].fill=PatternFill('solid',fgColor='FFF0AD')
    summary['B4'].comment=Comment('Criterio declarado de comparación numérica. OK cuando ABS(Obtenido − Esperado) ≤ esta tolerancia. No se redondean los datos antes de comparar.','SIDE')
    summary.merge_cells('A6:H6'); summary['A6']='Edite solo Obtenido (amarillo) al repetir los casos; Esperado conserva la evidencia del Excel original.'
    summary.merge_cells('A7:H7'); summary['A7']='NO COMPARABLE: falta de oráculo numérico o fórmula original con error. PENDIENTE: no existe observación SIDE.'
    headers=['Tema','Total','OK','DIFERENCIA','PENDIENTE','NO COMPARABLE','% OK comparables']
    for c,v in enumerate(headers,1): summary.cell(9,c,v)
    source_headers=['Caso','Tema','Resultado','Entradas Excel (editar solo copia)','Hoja','Celda','Fórmula original (texto)','Valor Excel original','Observación / alcance']
    for c,v in enumerate(source_headers,1): provenance.cell(7,c,v)
    counts={}; audit=[]
    for name,ws in sheets.items():
        title(ws,name.capitalize())
        ws.merge_cells('A4:J4'); ws['A4']="Tolerancia numérica: Resumen!B4. Diferencia = Obtenido − Esperado. Sin oráculo: NO COMPARABLE."
        ws.merge_cells('A5:J5'); ws['A5']='Entradas y fuentes completas en Fuentes originales. DOP se compara para la cantidad solicitada; el inventario es otro dato.'
        for c,v in enumerate(['Caso','Resultado / unidad','Entradas concretas SIDE','Fuente Excel','Esperado (Excel)','Obtenido (SIDE)','Diferencia','% diferencia','Estado','Fuente SIDE / pantalla'],1): ws.cell(7,c,v)
    for sc in source['escenarios']:
        ws=sheets[theme(sc['tema'])]
        for m in sc['mediciones']:
            record=index.get((sc['id'],m['key']),{})
            obtained=record.get('obtenido',record.get('valor'))
            expected=m.get('esperado'); comparable=m.get('comparable',True)
            numeric=isinstance(expected,(int,float)) and not isinstance(expected,bool)
            r=ws.max_row+1
            location=f"{m.get('hoja','')}!{m.get('celda','')}".strip('!')
            entries=sc.get('side',sc.get('entradas_side',{}))
            side_source=short(record.get('fuente',record.get('pantalla',record.get('source','Sin observación registrada'))))
            if record.get('screenshot'): side_source+=' | Captura: '+str(record['screenshot'])
            values=[sc['id'],m.get('nombre',m['key'])+(' ['+m['unidad']+']' if m.get('unidad') else ''),short(entries),location,expected,obtained,None,None,None,side_source]
            for c,v in enumerate(values,1):
                if isinstance(v,(dict,list)):v=short(v)
                ws.cell(r,c,v)
            if isinstance(expected,str) and expected.startswith('#'): ws.cell(r,5).data_type='s'
            ws.cell(r,6).fill=PatternFill('solid',fgColor='FFF0AD')
            ws.cell(r,6).font=Font(name='Arial',size=10,color='0000FF')
            if numeric and comparable:
                ws.cell(r,7,f'=IF(ISNUMBER(F{r}),F{r}-E{r},"")')
                ws.cell(r,8,f'=IF(ISNUMBER(F{r}),IF(E{r}=0,IF(F{r}=0,0,"N/A"),G{r}/ABS(E{r})),"")')
                ws.cell(r,9,f'=IF(ISNUMBER(F{r}),IF(ABS(G{r})<=Resumen!$B$4,"OK","DIFERENCIA"),"PENDIENTE")')
            elif comparable:
                ws.cell(r,9,f'=IF(ISBLANK(F{r}),"PENDIENTE",IF(F{r}=E{r},"OK","DIFERENCIA"))')
            else:
                ws.cell(r,9,'NO COMPARABLE')
                ws.cell(r,7,None);ws.cell(r,8,None)
            for c in [5,6,7]:ws.cell(r,c).number_format='#,##0.######;[Red](#,##0.######);0'
            ws.cell(r,8).number_format='0.0000%'
            ws.row_dimensions[r].height=62
            pr=provenance.max_row+1
            pvals=[sc['id'],sc['tema'],m.get('nombre',m['key']),short(sc.get('entradas_excel',{})),m.get('hoja'),m.get('celda'),m.get('formula'),expected,m.get('nota',sc.get('nota',short(sc.get('notas',''))))]
            for c,v in enumerate(pvals,1):
                provenance.cell(pr,c,short(v) if isinstance(v,(dict,list)) else v)
                if isinstance(v,str) and v.startswith(('=','#')):provenance.cell(pr,c).data_type='s'
            audit.append({'caso':sc['id'],'key':m['key'],'sheet':ws.title,'row':r,'esperado':expected,'obtenido':obtained,'comparable':comparable})
    for n,(name,ws) in enumerate(sheets.items(),10):
        last=ws.max_row; counts[name]=max(0,last-7)
        summary.cell(n,1,name);summary.cell(n,2,f'=COUNTA(\'{name}\'!A8:A{max(8,last)})')
        for col,state in [(3,'OK'),(4,'DIFERENCIA'),(5,'PENDIENTE'),(6,'NO COMPARABLE')]:summary.cell(n,col,f'=COUNTIF(\'{name}\'!I8:I{max(8,last)},"{state}")')
        summary.cell(n,7,f'=IFERROR(C{n}/(C{n}+D{n}),0)').number_format='0.0%'
        for state,color in [('OK','DDF1E4'),('DIFERENCIA','FADDDD'),('PENDIENTE','FFF0AD'),('NO COMPARABLE','E6EAED')]:
            ws.conditional_formatting.add(f'A8:J{max(8,last)}',FormulaRule(formula=[f'$I8="{state}"'],fill=PatternFill('solid',fgColor=color)))
        ws.auto_filter.ref=f'A7:J{max(8,last)}';ws.print_title_rows='1:7'
        style(ws,dict(A=10,B=31,C=48,D=35,E=20,F=20,G=18,H=18,I=22,J=48))
        for row in ws.iter_rows(min_row=8):
            for c in [row[4],row[5]]:c.font=Font(name='Arial',size=10,color='0000FF')
            row[8].font=Font(name='Arial',size=10,color='008000')
        title(ws,name.capitalize())
    summary.cell(15,1,'TOTAL')
    for c in range(2,7):summary.cell(15,c,f'=SUM({chr(64+c)}10:{chr(64+c)}14)')
    summary.cell(15,7,'=IFERROR(C15/(C15+D15),0)').number_format='0.0%'
    notes=[('Escenarios únicos',len(source['escenarios'])),('Fuente de esperados',source.get('archivo',source.get('fuente','DECISIONES_SIDE_16.xlsx'))),('Celdas amarillas','Obtenido: escriba lo observado en SIDE. B4: tolerancia elegida.'),('Diferencia porcentual','(Obtenido − Esperado) / ABS(Esperado). Si Esperado = 0 y Obtenido ≠ 0: N/A.'),('Estado OK','No equivale a auditoría del juego completo. Se limita al resultado y entradas de la fila.'),('Errores del original','Un error como #DIV/0! se conserva como texto para evidenciarlo; no se transforma en cero.'),('Caja y faltante','Las reglas del juego sin celda equivalente en el original quedan NO COMPARABLE.'),('Cómo añadir casos','Duplique una fila del tema; anote entradas y celda fuente. Amplíe filtros y rangos del resumen.')]
    for r,(a,b) in enumerate(notes,18): summary.cell(r,1,a);summary.merge_cells(start_row=r,start_column=2,end_row=r,end_column=8);summary.cell(r,2,short(b));summary.row_dimensions[r].height=30
    style(summary,dict(A=29,B=18,C=15,D=18,E=18,F=23,G=24,H=14));title(summary,'Verificación manual · SIDE vs Excel','H');summary.freeze_panes='B10'
    style(provenance,dict(A=10,B=24,C=30,D=60,E=30,F=15,G=95,H=23,I=65));title(provenance,'Trazabilidad · Excel original','I');provenance.auto_filter.ref=f'A7:I{provenance.max_row}'
    for ws in wb:
        hr=9 if ws==summary else 7
        for c in ws[hr]:c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10);c.fill=PatternFill('solid',fgColor=TEAL)
        ws.row_dimensions[hr].height=32
    original_errors=wb.create_sheet('Errores del original')
    title(original_errors,'Errores heredados · evidencia','E')
    original_errors.merge_cells('A4:E5');original_errors['A4']='Estos errores pertenecen al Excel original recalculado. Son evidencia, no fórmulas con error del comparativo. Se excluyen del conteo de igualdad numérica.'
    for c,v in enumerate(['Caso','Hoja','Celda','Error original (texto)','Interpretación'],1):original_errors.cell(7,c,v)
    for sc in source['escenarios']:
        for addr,error in sc.get('distribucion_no_comparable',{}).items():
            row=original_errors.max_row+1
            for c,v in enumerate([sc['id'],'CANTIDAD DE PRODUCCIÓN',addr,error,'Plan cero: la proporción divide entre cero; comparar necesidades brutas en la hoja DOP.'],1):
                original_errors.cell(row,c,v)
                if c==4:original_errors.cell(row,c).data_type='s'
    style(original_errors,dict(A=12,B=34,C=14,D=25,E=90));title(original_errors,'Errores heredados · evidencia','E')
    for c in original_errors[7]:c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10);c.fill=PatternFill('solid',fgColor=TEAL)
    original_errors.freeze_panes='A8'
    # Normalize package metadata before delivery. Native Excel rejects the
    # original combination of VML comment/print metadata although openpyxl
    # reads it. A fresh workbook preserves data, formulas and visible styles.
    normalized=Workbook();normalized.remove(normalized.active)
    for sheet in wb:
        clean=normalized.create_sheet(sheet.title)
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    c=clean.cell(cell.row,cell.column,cell.value);c.data_type=cell.data_type
                    c.font=copy(cell.font);c.fill=copy(cell.fill);c.border=copy(cell.border)
                    c.alignment=copy(cell.alignment);c.number_format=cell.number_format
        for key,dim in sheet.column_dimensions.items():clean.column_dimensions[key].width=dim.width
        for key,dim in sheet.row_dimensions.items():clean.row_dimensions[key].height=dim.height
        for merged in sheet.merged_cells.ranges:clean.merge_cells(str(merged))
        clean.conditional_formatting=copy(sheet.conditional_formatting)
        clean.freeze_panes=sheet.freeze_panes;clean.auto_filter.ref=sheet.auto_filter.ref
        clean.sheet_view.showGridLines=False
    path=OUT/'COMPARATIVO_SIDE_vs_EXCEL.xlsx';normalized.save(path)
    (OUT/'auditoria_comparativo.json').write_text(json.dumps({'archivo':path.name,'escenarios':len(source['escenarios']),'filas_por_tema':counts,'filas':audit},ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'archivo':str(path),'escenarios':len(source['escenarios']),'filas_por_tema':counts},ensure_ascii=False))

if __name__=='__main__':build()

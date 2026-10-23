"""Independent read-only verification of original Excel, cases and deliverables."""
from pathlib import Path
import json, hashlib, math
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'reports/verificacion'
def read(name): return json.loads((OUT/name).read_text(encoding='utf-8-sig'))
def equal(a,b):
    if isinstance(a,(int,float)) and isinstance(b,(int,float)): return math.isclose(a,b,rel_tol=0,abs_tol=1e-10)
    return a==b
def norm_formula(value):
    if hasattr(value,'text'):value=value.text
    return value.upper().replace('=+','=').replace('_XLUDF.','').replace('_XLFN.','') if isinstance(value,str) else value
def audit():
    x=read('datos_excel.json'); s=read('datos_side.json')
    errors=[]; checks=[]
    def check(name,condition,detail=None):
        checks.append({'control':name,'ok':bool(condition),'detalle':detail})
        if not condition: errors.append({'control':name,'detalle':detail})
    source=ROOT/x['fuente']
    check('SHA256 del Excel original',hashlib.sha256(source.read_bytes()).hexdigest()==x['sha256'],x['sha256'])
    original=openpyxl.load_workbook(source,data_only=False)
    side={c['id']:c for c in s['escenarios']}
    metrics=[]
    for case in x['escenarios']:
        sid=case['id']; path=ROOT/case['copia_recalculada'] if case.get('copia_recalculada') else None
        form=openpyxl.load_workbook(path,data_only=False) if path else None
        value=openpyxl.load_workbook(path,data_only=True) if path else None
        sm={m['key']:m for m in side[sid]['mediciones']}
        if form:
            for sheet,inputs in case['entradas_excel'].items():
                for addr,v in inputs.items():
                    check(f'{sid} entrada {sheet}!{addr}',equal(form[sheet][addr].value,v),{'esperado':v,'copia':form[sheet][addr].value})
            for sh in original:
                for row in sh:
                    for c in row:
                        if c.data_type=='f':
                            # Excel COM may store leading _xlfn aliases; original common formulas stay literal.
                            check(f'{sid} fórmula preservada {sh.title}!{c.coordinate}',norm_formula(form[sh.title][c.coordinate].value)==norm_formula(c.value))
        for m in case['mediciones']:
            expected=m.get('esperado'); actual=sm.get(m['key'],{}).get('obtenido')
            if value:
                raw=value[m['hoja']][m['celda']].value
                check(f'{sid} esperado desde copia Excel {m["key"]}',equal(raw,expected),{'excel':raw,'json':expected})
                check(f'{sid} fórmula documentada {m["key"]}',norm_formula(original[m['hoja']][m['celda']].value)==norm_formula(m.get('formula')))
            if m.get('comparable',True):
                check(f'{sid} Excel vs SIDE {m["key"]}',isinstance(actual,(int,float)) and abs(actual-expected)<=x['tolerancia_absoluta'],{'excel':expected,'side':actual})
            metrics.append({'id':sid,'key':m['key'],'esperado':expected,'obtenido':actual,'comparable':m.get('comparable',True)})
        if form: form.close();value.close()
    rules=original['Matriz_Consecuencias']
    check('Regla original frontera Premium',rules['C2'].value=='< 4 y Segmento = Premium' and rules['C3'].value=='>= 7 y Segmento = Premium',{'C2':rules['C2'].value,'D2':rules['D2'].value,'C3':rules['C3'].value,'D3':rules['D3'].value})
    for sid,expected_pair in [('E14',[1,0]),('E15',[1.4,2])]:
        case=next(c for c in x['escenarios'] if c['id']==sid)
        check(f'{sid} interpretación de regla textual', [m['esperado'] for m in case['mediciones']]==expected_pair)
    evidence=ROOT/'tests/output/verificacion'
    browser=json.loads((evidence/'browser-escenarios.json').read_text(encoding='utf8'))
    components=json.loads((evidence/'browser-componentes.json').read_text(encoding='utf8'))
    bi={c['id']:c for c in browser['scenarios']};bi.update({c['id']:c for c in components})
    for m in metrics:
        if m['comparable']:
            recorded=bi[m['id']]['values'].get(m['key'])
            check(f'{m["id"]} navegador vs obtenido {m["key"]}',equal(recorded,m['obtenido']),{'browser':recorded,'side':m['obtenido']})
    for name,needle in [('sim-1.log','PASS SIM-1'),('sim-2.log','PASS SIM-2'),('sim-3.log','PASS SIM-3'),('node.log','fail 0'),('browser-dop.log','PASS no uncaught browser errors'),('browser-costos.log','PASS 153 independent checks'),('browser-ciclo.log','PASS page scroll')]:
        check(f'Evidencia prueba {name}',needle in (evidence/name).read_text(encoding='utf8'))
    oracle=json.loads((evidence/'oraculo.log').read_text(encoding='utf8'))
    check('Oráculo Excel original cero diferencias',oracle['differences']==[] and oracle['source_sha256']==x['sha256'],{'comparaciones':oracle['comparisons'],'numericas':oracle['numeric_comparisons'],'errores_originales':oracle['source_errors']})
    wbpath=OUT/'COMPARATIVO_SIDE_vs_EXCEL.xlsx'
    if wbpath.exists():
        f=openpyxl.load_workbook(wbpath,data_only=False);v=openpyxl.load_workbook(wbpath,data_only=True)
        check('Tolerancia visible Resumen!B4',f['Resumen']['B4'].value==0.01)
        themes=['DOP por cantidad','costo de producción','caja','consecuencias','puntajes']
        seen=[]
        for theme in themes:
            check(f'Hoja requerida {theme}',theme in f.sheetnames)
            sh=f[theme]; cv=v[theme]
            check(f'Formato condicional verde/rojo {theme}',len(sh.conditional_formatting)>=1)
            for r in range(8,sh.max_row+1):
                sid=sh.cell(r,1).value
                if not sid:continue
                expected=sh.cell(r,5).value;actual=sh.cell(r,6).value
                loc=sh.cell(r,4).value
                candidates=[m for c in x['escenarios'] if c['id']==sid for m in c['mediciones'] if (f"{m.get('hoja','')}!{m.get('celda','')}".strip('!')==loc)]
                check(f'{theme} fila{r} esperado trazable',any(equal(expected,m.get('esperado')) for m in candidates))
                candidates2=[m for m in metrics if m['id']==sid and equal(m['esperado'],expected) and equal(m['obtenido'],actual)]
                check(f'{theme} fila{r} obtenido trazable',bool(candidates2))
                if isinstance(expected,(int,float)):
                    check(f'{theme} fila{r} diferencia fórmula',sh.cell(r,7).data_type=='f' and f'F{r}-E{r}' in sh.cell(r,7).value)
                    check(f'{theme} fila{r} diferencia cache',equal(cv.cell(r,7).value,actual-expected))
                    check(f'{theme} fila{r} porcentaje fórmula',sh.cell(r,8).data_type=='f')
                    check(f'{theme} fila{r} estado fórmula',sh.cell(r,9).data_type=='f')
                    check(f'{theme} fila{r} estado cache',cv.cell(r,9).value=='OK')
                else:check(f'{theme} fila{r} sin oráculo',sh.cell(r,9).value=='NO COMPARABLE')
                seen.append(sid)
        check('112 registros workbook',len(seen)==112,len(seen))
        check('15 escenarios workbook',len(set(seen))==15)
        check('Conteos Resumen cache', [v['Resumen'].cell(15,c).value for c in range(2,7)]==[112,111,0,0,1],[v['Resumen'].cell(15,c).value for c in range(2,7)])
        check('Workbook sin errores fórmula',not [(sh.title,c.coordinate,c.value) for sh in v for row in sh for c in row if c.data_type=='e'])
        f.close();v.close()
    pdfpath=OUT/'GUIA_VERIFICACION_SIDE_vs_EXCEL.pdf'
    if pdfpath.exists():
        import pypdfium2
        pdf=pypdfium2.PdfDocument(str(pdfpath))
        pages=[p.get_textpage().get_text_range() for p in pdf]
        check('PDF 27 páginas A4',len(pdf)==27 and all(abs(pdf[i].get_width()-595.276)<1 and abs(pdf[i].get_height()-841.89)<1 for i in range(len(pdf))))
        for case in x['escenarios']:
            sid=case['id'];matches=[p for p in pages if f'{sid} · ' in p and 'ESCENARIO REPRODUCIBLE / ALCANCE DECLARADO' in p]
            check(f'PDF sección individual {sid}',len(matches)==1)
            if matches:
                text=matches[0]
                check(f'PDF tolerancia {sid}','0.01' in text and 'Resumen!B4' in text)
                for m in case['mediciones']:
                    if m['comparable']:
                        displayed=f'{m["esperado"]:,.6f}'.rstrip('0').rstrip('.')
                        check(f'PDF esperado {sid} {m["key"]}',displayed in text,displayed)
        qa=json.loads((OUT/'evidencia/revision_visual_pdf.json').read_text(encoding='utf8'))
        check('Revisión visual documentada página por página',qa['paginas_renderizadas']==27 and qa['paginas_revisadas_visualmente']==list(range(1,28)))
        pdf.close()
    original.close()
    result={'rol':'Auditor independiente','original_sha256':x['sha256'],'escenarios':len(x['escenarios']),'metricas_comparables':sum(m['comparable'] for m in metrics),'no_comparables':sum(not m['comparable'] for m in metrics),'controles':len(checks),'fallos':len(errors),'errores':errors,'detalle':checks}
    (OUT/'auditoria_independiente.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({k:v for k,v in result.items() if k not in ('detalle','errores')},ensure_ascii=False));print(json.dumps(errors[:10],ensure_ascii=False))
if __name__=='__main__':audit()

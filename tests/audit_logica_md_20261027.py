"""Extract read-only workbook evidence for the Markdown conformity audit."""
import hashlib
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path('C:/Users/Asus/Downloads/DECISIONES SIDE (19).xlsx')
MD = Path('C:/Users/Asus/Downloads/Logica_SIDE_Costos_de_Produccion.md')
OUT = ROOT / 'tests/output/auditoria_md_20261027'
OUT.mkdir(parents=True, exist_ok=True)

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

book = openpyxl.load_workbook(SOURCE, data_only=False)
cached = openpyxl.load_workbook(SOURCE, data_only=True)
sheet = book['Costos De Producción']
cells = {}
for row in sheet.iter_rows(min_row=1, max_row=152, min_col=1, max_col=30):
    for cell in row:
        if cell.value is not None:
            cells[cell.coordinate] = {'formula_or_value': cell.value,
                                      'cached_value': cached[sheet.title][cell.coordinate].value}

repo = openpyxl.load_workbook(ROOT / 'docs/DECISIONES_SIDE_19.xlsx', data_only=False)
other = repo[sheet.title]
differences = []
for row in sheet.iter_rows(min_row=1, max_row=152, min_col=1, max_col=30):
    for cell in row:
        if cell.value != other[cell.coordinate].value:
            differences.append({'cell': cell.coordinate, 'attached': cell.value,
                                'repository': other[cell.coordinate].value})
copy = book['Copia de Costos De Producción']
copy_differences = [cell.coordinate for row in sheet.iter_rows() for cell in row
                    if cell.value != copy[cell.coordinate].value]

production_sheet = book[book.sheetnames[5]]
production_drivers = {key: production_sheet[key].value for key in
                      ['E5','E6','E7','E11','E12','E13','E17','E18','E19','L5','M5']}
# MD 3.1 / 3.2 corrected example: actual production, rather than desired output,
# feeds the independent sequential consumption and valuation calculation.
available = [60,110,30,200,190,0,1500,1400]
prices = [25,45,80,4,9,2,4,7]
used = [0] * 8
premium_ratios = []
for indices, needed in zip([[0,1,2],[3,4],[5,6,7]], [155.6,368,2894]):
    for index in indices:
        used[index] = min(available[index], needed)
        needed -= used[index]
    premium_ratios.append(used[indices[-1]] / sum(used[index] for index in indices))
mp_cost = sum(q * p * .92 for q, p in zip(used, prices))
corrected_example = {
    'production': [150,140,78], 'actual_consumption': used,
    'material_cost': mp_cost,
    'closing_material_value': sum(q*p*.92 for q,p in zip(available,prices))-mp_cost,
    'reputation_0_10': (78/368 + 2/7 + sum(premium_ratios)/3 + .3)*2.5,
}

evidence = {
    'source': str(SOURCE), 'source_sha256': digest(SOURCE),
    'markdown': str(MD), 'markdown_sha256': digest(MD),
    'repository_workbook_sha256': digest(ROOT / 'docs/DECISIONES_SIDE_19.xlsx'),
    'sheet': sheet.title, 'sheets': book.sheetnames, 'cells': cells,
    'attached_vs_repository_cost_sheet_differences': differences,
    'copy_sheet_value_formula_differences': copy_differences,
    'other_production_sheet_drivers': production_drivers,
    'corrected_md_example': corrected_example,
    'method': 'openpyxl read only inspection in memory; cached values are not a fresh Excel recalculation',
}
(OUT / 'excel_evidence.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
keys = ['E5','E6','E7','E11','E12','E13','E17','E18','E19','L5','M5','N5','E46','E47','E48','H56','H58','H61','C113','D113','E113','F113','G113','AD96','W95','W100']
print(json.dumps({'source_sha256': digest(SOURCE), 'key_cells': {key: cells.get(key) for key in keys},
                  'repository_cost_sheet_difference_count': len(differences), 'copy_sheet_differences': copy_differences},
                 ensure_ascii=False, indent=2, default=str))

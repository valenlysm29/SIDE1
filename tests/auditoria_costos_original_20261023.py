"""Extract original Excel cells, formulas and cached values without saving the workbook.

Run from SIDE1 with the bundled Python (openpyxl). This is an extraction,
not a recalculation: the fixture explicitly records that distinction.
"""
from pathlib import Path
import hashlib
import json
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / '.codex-work/DECISIONES_SIDE_16.xlsx'
DEST = ROOT / 'tests/oracle/auditoria_costos_original_20261023.json'


def main():
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    assert digest == '3b6b42bbbc56a2534ce7e241c607d572dd38dcf02311683609497df36a0c1eb1'
    formulas = openpyxl.load_workbook(SOURCE, data_only=False)
    values = openpyxl.load_workbook(SOURCE, data_only=True)
    sheet = formulas.sheetnames[6]

    def cell(address):
        return {'sheet': sheet, 'cell': address,
                'formula_or_literal': formulas[sheet][address].value,
                'value': values[sheet][address].value}

    labor = []
    fields = ['base', 'essalud', 'gratificacion', 'cts', 'vacaciones', 'sumaAnual', 'pagoMensual']
    for role, columns in [('corte', ['U', 'V']), ('ensamble', ['Y', 'Z']), ('acabado', ['AC', 'AD'])]:
        for level, start in enumerate([5, 14, 23], 1):
            for regime, col in zip(['MYPE', 'NO_MYPE'], columns):
                labor.append({'role': role, 'level': level, 'regime': regime,
                              'expected': {field: cell(f'{col}{start + offset}') for offset, field in enumerate(fields)}})
    for role, columns in [('jefatura', ['U', 'V']), ('limpieza', ['Y', 'Z'])]:
        for regime, col in zip(['MYPE', 'NO_MYPE'], columns):
            labor.append({'role': role, 'level': 1, 'regime': regime,
                          'expected': {field: cell(f'{col}{33 + offset}') for offset, field in enumerate(fields)}})
    assets = []
    for code, cols, start, mold in [('mesa', ['U', 'V'], 45, False), ('INF-02', ['Y', 'Z'], 45, False),
                                  ('INF-03', ['Y', 'Z'], 53, False), ('INF-04', ['Y', 'Z'], 61, False),
                                  ('INF-05', ['AC', 'AD'], 45, False), ('INF-06', ['AC', 'AD'], 53, False),
                                  ('INF-07', ['AC', 'AD'], 61, False), ('INF-08', ['U', 'V'], 53, True),
                                  ('INF-09', ['U', 'V'], 60, True), ('INF-10', ['U', 'V'], 67, True)]:
        fields_asset = ['adquisicion', 'capacitacion', 'mantenimientoAnual', 'depreciacionAnual', 'sumaAnual'] if mold else ['adquisicion', 'instalacion', 'capacitacion', 'mantenimientoAnual', 'depreciacionAnual', 'sumaAnual']
        for regime, col in zip(['MYPE', 'NO_MYPE'], cols):
            assets.append({'code': code, 'regime': regime,
                           'expected': {field: cell(f'{col}{start + offset}') for offset, field in enumerate(fields_asset)}})
    materials = [{'code': f'INS-{i+1:02}', 'price': cell(f'V{76+i}'),
                  'discounted_unit_price': cell(f'W{76+i}'), 'conversion': cell(f'{col}32')}
                 for i, col in enumerate('DEFGHIJK')]
    result = {'provenance': {'source': str(SOURCE.relative_to(ROOT)), 'sha256': digest,
                           'method': 'openpyxl dual-load original formulas + original cached values; no recalculation',
                           'generator': str(Path(__file__).relative_to(ROOT))},
              'labor': labor, 'assets': assets, 'materials': materials,
              'local': [cell('AC33'), cell('AD33')]}
    DEST.write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    formulas.close()
    values.close()
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == digest
    print(json.dumps({'labor_profiles': len(labor), 'asset_profiles': len(assets), 'materials': len(materials),
                      'original_cells': sum(len(x['expected']) for x in labor + assets) + len(materials)*3 + 2}))


if __name__ == '__main__':
    main()

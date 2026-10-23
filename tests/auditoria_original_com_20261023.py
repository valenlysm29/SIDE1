"""Recalculate only copies of the original workbook; no model extensions."""
import hashlib
import json
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / 'tests/oracle/auditoria_original_20261023_planes.json'
OUT = ROOT / 'tests/output/auditoria-20261023/original-com'

def main():
    plan = json.loads(MANIFEST.read_text(encoding='utf-8'))
    source = ROOT / plan['source']
    assert hashlib.sha256(source.read_bytes()).hexdigest() == plan['sha256']
    OUT.mkdir(parents=True, exist_ok=True)
    all_plans = plan['planes'] + plan.get('planes_operando_inyectado', [])
    batch_manifest = OUT / 'manifest_batch.json'
    if '--read-existing' in sys.argv:
        executed = json.loads(batch_manifest.read_text(encoding='utf-8'))
        assert executed['sha256'] == plan['sha256']
        assert {row['nombre']: row['entradas'] for row in executed['planes']} == {row['nombre']: row['entradas'] for row in all_plans}, 'Los operandos deben coincidir con el recálculo existente'
    else:
        batch_manifest.write_text(json.dumps({**plan, 'planes': all_plans}, ensure_ascii=False), encoding='utf-8')
        subprocess.run(['powershell.exe', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File',
                        str(ROOT / 'tests/oracle/recalcular_original_lote_20261023.ps1'),
                        '-Manifest', str(batch_manifest), '-Destination', str(OUT)], cwd=ROOT, check=True)
    metadata = {row['nombre']: row for row in json.loads((OUT / 'metadata.json').read_text(encoding='utf-8-sig'))}
    original = openpyxl.load_workbook(source, read_only=True)
    sheets = original.sheetnames
    original.close()
    cases = []
    for case in plan['planes']:
        book = openpyxl.load_workbook(OUT / (case['nombre'] + '.xlsx'), data_only=True)
        assert book.sheetnames == sheets, 'No se permiten hojas añadidas'
        values = {key: book[sheet][cell].value for key, (sheet, cell) in case['salidas'].items()}
        errors = [{'sheet': sheet.title, 'cell': cell.coordinate, 'error': cell.value}
                  for sheet in book for row in sheet for cell in row
                  if cell.data_type == 'e' or isinstance(cell.value, str) and cell.value.startswith(('#DIV/0!', '#NAME?', '#VALUE!', '#REF!'))]
        cases.append({'name': case['nombre'], 'metadata': metadata[case['nombre']], 'values': values, 'source_errors': errors})
        book.close()
    operand_experiments = []
    for case in plan.get('planes_operando_inyectado', []):
        book = openpyxl.load_workbook(OUT / (case['nombre'] + '.xlsx'), data_only=True)
        values = {key: book[sheet][cell].value for key, (sheet, cell) in case['salidas'].items()}
        operand_experiments.append({'name': case['nombre'], 'metadata': metadata[case['nombre']],
                                    'formula_modified_in_copy': True, 'values': values})
        book.close()
    fixture = {'source': plan['source'], 'source_sha256': plan['sha256'], 'engine': 'Microsoft Excel COM',
               'generated_at': datetime.now(timezone.utc).isoformat(), 'added_sheets': 0, 'cases': cases,
               'operand_experiments': operand_experiments}
    dest = ROOT / 'tests/oracle/auditoria_original_recalculado_20261023.json'
    dest.write_text(json.dumps(fixture, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    assert hashlib.sha256(source.read_bytes()).hexdigest() == plan['sha256']
    print(json.dumps({'cases': len(cases), 'original_formula_exports': sum(len(row['values']) for row in cases), 'source_preserved': True}), flush=True)

if __name__ == '__main__':
    main()

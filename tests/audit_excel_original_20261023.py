"""Read-only audit of the untouched source; no COM, no workbook writes.

Run from SIDE1 with Python containing openpyxl. The report explicitly separates
formula inventory, original cached-value comparisons, and rules absent from Excel.
"""
import hashlib
import json
import math
import re
import shutil
import subprocess
from collections import Counter, defaultdict
from pathlib import Path
import openpyxl
from openpyxl.formula.translate import Translator

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / '.codex-work/DECISIONES_SIDE_16.xlsx'
SHA = '3b6b42bbbc56a2534ce7e241c607d572dd38dcf02311683609497df36a0c1eb1'

def main():
    before = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    assert before == SHA
    formulas = openpyxl.load_workbook(SOURCE, data_only=False)
    cache = openpyxl.load_workbook(SOURCE, data_only=True)
    quantity = next(s for s in formulas.sheetnames if s.startswith('CANTIDAD'))
    costs = next(s for s in formulas.sheetnames if s.startswith('Costos De'))
    build_com_plans(formulas,quantity,costs)
    node = shutil.which('node') or str(Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe')
    js = "const c=require('./js/engine/costos.js'); console.log(JSON.stringify({config:require('./js/engine/config_excel.js').CONFIG,catalogo:require('./js/engine/catalogo_excel.js'),labor:Object.fromEntries(['corte','ensamble','acabado','jefatura','limpieza'].flatMap(r=>[1,2,3].flatMap(n=>['MYPE','NO_MYPE'].map(t=>[r+'_'+n+'_'+t,c.costoLaboral(r,n,t)])))),assets:Object.fromEntries(['mesa','INF-02','INF-03','INF-04','INF-05','INF-06','INF-07','INF-08','INF-09','INF-10'].flatMap(k=>['MYPE','NO_MYPE'].map(t=>[k+'_'+t,c.costoMaquinaria(k,t)])))}));"
    side = json.loads(subprocess.check_output([node,'-e',js],cwd=ROOT,text=True,encoding='utf-8'))
    all_cells=[]; families=defaultdict(list); counts=[]
    for sheet in formulas:
        counter=Counter()
        for row in sheet:
            for c in row:
                if c.value is None: continue
                kind='formula' if c.data_type=='f' else 'numeric_literal' if isinstance(c.value,(int,float)) else 'text'
                counter[kind]+=1
                literal=c.value.text if hasattr(c.value,'text') else c.value
                item={'sheet':sheet.title,'cell':c.coordinate,'kind':kind,'original':literal,'cached':cache[sheet.title][c.coordinate].value}
                all_cells.append(item)
                if kind=='formula':
                    try: normalized=Translator(literal,origin=c.coordinate).translate_formula('CV1000')
                    except Exception: normalized=literal
                    key=(sheet.title,normalized)
                    families[key].append(c.coordinate)
        counts.append({'sheet':sheet.title,**counter})
    comparisons=[]
    def check(label,sheet,cell,actual,source_path):
        expected=cache[sheet][cell].value
        equal=math.isclose(expected,actual,rel_tol=1e-10,abs_tol=1e-9) if isinstance(expected,(int,float)) and isinstance(actual,(int,float)) else expected==actual
        comparisons.append({'label':label,'excel_sheet':sheet,'excel_cell':cell,'excel_formula_or_literal':formulas[sheet][cell].value,'excel_cached':expected,'side':actual,'side_source':source_path,'equal':equal})
    config=side['config']
    for area,start in [('corte',5),('ensamble',11),('acabado',17)]:
        for i in range(3): check('productivity '+area,quantity,'E'+str(start+i),config['PRODUCTIVIDAD'][area][i],'js/engine/config_excel.js PRODUCTIVIDAD')
    for i,code in enumerate(['INS-01','INS-02','INS-03','INS-04','INS-05','INS-06','INS-07','INS-08']):
        check('material price '+code,costs,'V'+str(76+i),config['PRECIO_MP'][code],'js/engine/config_excel.js PRECIO_MP')
        check('purchase conversion '+code,costs,openpyxl.utils.get_column_letter(4+i)+'32',config['CONVERSION_COMPRA'][code],'js/engine/config_excel.js CONVERSION_COMPRA')
    for code,row,col in [('INF-08',31,4),('INF-09',32,5),('INF-10',33,6)]:
        for family,delta in [('cuero',0),('accesorios',3),('hilo',6)]: check('BOM '+code+' '+family,quantity,openpyxl.utils.get_column_letter(col+delta)+str(row),config['CONSUMO_MOLDE'][code][family],'js/engine/config_excel.js CONSUMO_MOLDE')
    for role,col in [('corte','U'),('ensamble','Y'),('acabado','AC'),('jefatura','U'),('limpieza','Y')]:
        for level in ([1,2,3] if role in ['corte','ensamble','acabado'] else [1]):
            start={1:5,2:14,3:23}[level] if role in ['corte','ensamble','acabado'] else 33
            for regime in ['MYPE','NO_MYPE']:
                cl=col if regime=='MYPE' else openpyxl.utils.get_column_letter(openpyxl.utils.column_index_from_string(col)+1)
                obj=side['labor'][role+'_'+str(level)+'_'+regime]
                for field,offset in [('base',0),('essalud',1),('gratificacion',2),('cts',3),('vacaciones',4),('sumaAnual',5),('pagoMensual',6)]: check(role+str(level)+' '+regime+' '+field,costs,cl+str(start+offset),obj[field],'js/engine/costos.js costoLaboral')
    assets=[('mesa','U',45,False),('INF-02','Y',45,False),('INF-03','Y',53,False),('INF-04','Y',61,False),('INF-05','AC',45,False),('INF-06','AC',53,False),('INF-07','AC',61,False),('INF-08','U',53,True),('INF-09','U',60,True),('INF-10','U',67,True)]
    for code,col,row,mold in assets:
        for regime in ['MYPE','NO_MYPE']:
            cl=col if regime=='MYPE' else openpyxl.utils.get_column_letter(openpyxl.utils.column_index_from_string(col)+1)
            obj=side['assets'][code+'_'+regime]
            fields=[('adquisicion',0),('capacitacion',1),('mantenimientoAnual',2),('depreciacionAnual',3),('sumaAnual',4)] if mold else [('adquisicion',0),('instalacion',1),('capacitacion',2),('mantenimientoAnual',3),('depreciacionAnual',4),('sumaAnual',5)]
            for field,offset in fields: check(code+' '+regime+' '+field,costs,cl+str(row+offset),obj[field],'js/engine/costos.js costoMaquinaria')
    for code,units,penetration,weights,ideal in [('SEG-01','B89','B90',['M124','M125','M126'],None),('SEG-02','B103','B104',['M137','M138','M139'],'M140'),('SEG-03','B117','B118',['M150','M151','M152'],'M153')]:
        market=config['MERCADOS'][code]
        check(code+' units',quantity,units,market['unidades'],'js/engine/config_excel.js MERCADOS')
        check(code+' penetration',quantity,penetration,market['penetracion'],'js/engine/config_excel.js MERCADOS')
        for i,a in enumerate(weights): check(code+' weight '+str(i),quantity,a,market['pesos'][i],'js/engine/config_excel.js MERCADOS')
        if ideal: check(code+' ideal',quantity,ideal,market['precioIdeal'],'js/engine/config_excel.js MERCADOS')
    for row in side['catalogo']:
        for key,column in [('codigo','A'),('costo_original','G'),('calidad','H'),('eficiencia','I'),('reputacion','J'),('financiero','K')]:
            check('catalog '+row['codigo']+' '+key,'Catalogo_Decisiones',column+str(row['fila_excel']),row.get(key),'js/engine/catalogo_excel.js')
    contradictions=[
        {'id':'capacity_sheets','status':'contradiction','cells':[[quantity,'L5'],[costs,'L5'],[quantity,'E5:E7'],[costs,'E5:E7'],[quantity,'E17:E19'],[costs,'E17:E19']],'detail':'24 versus 30 days; cutting 5/7/9 versus 6/8/12; finishing 6/7/10 versus 6/8/12.'},
        {'id':'market_model1_medium_reference','status':'contradiction','cells':[[quantity,'H97:H100'],[quantity,'B91'],[quantity,'B105']],'detail':'Medium segment model1 references economic market B91, not its B105. Both original caches equal 3500; baseline equality hides dependency error.'},
        {'id':'market_model2_quality_reference','status':'contradiction','cells':[[quantity,'E141:E144'],[quantity,'E154:E157'],[quantity,'D132'],[quantity,'D145'],[quantity,'D158']],'detail':'Medium and premium model2 quality percentages divide by economic reputation total D132, not their own totals. Default company reputations are equal across examples, hiding dependency difference.'},
        {'id':'zero_quantity','status':'formula_domain_error','cells':[[quantity,'E37:E39']],'detail':'Zero total quantity causes division by zero; literal original contains no guard.'},
        {'id':'allocation_overplan','status':'formula_behavior','cells':[[quantity,'E37:E39'],[quantity,'D56:E58']],'detail':'Capacity allocation can exceed an individual desired plan; later MIN limits actual output. ROUNDDOWN allocations may leave capacity unallocated.'},
        {'id':'thread_purchase_cost','status':'contradiction','cells':[[costs,'J31:J33'],[costs,'AB82:AD82'],[costs,'V93']],'detail':'6 reels convert to 1500 m but purchase example AB82 values only 6 m; inventory values 1500 m.'}]
    absent=[{'rule':r,'status':'absent_or_not_uniquely_defined','detail':d} for r,d in [
        ('inverse_staff_machines_time','No inverse planner/optimization/rounding recipe from quantity to staff, machines or lead time.'),
        ('overtime_tariff','CONS-06 describes overtime but supplies no amount or percentage.'),
        ('storage_tariff','CONS-07 describes storage but supplies no monetary rate.'),
        ('additional_bank_rate','CONS-05 increases next loan rate without number of percentage points.'),
        ('inventory_physical_merma','No universal material-waste percentage; explicit BOM and efficiency/capacity equations must remain distinct.'),
        ('full_cash_transition','Original has no complete per-cycle cash ledger integrating loans, sales collection schedules, salaries, acquisitions and inventory; game transition needs declared extensions.'),
        ('depreciation_units_lifetime','Catalog names units-of-production method but no lifetime output of 100000 units is defined.'),
        ('market_tie_guard','Economic price min=max division is unguarded; game tie fallback is an extension.'),
        ('negative_price_score_guard','Price-ideal formula can be negative; original has no clamp in that cell.')]]
    consequence_rows=[]
    for row in range(2,10): consequence_rows.append({'id':formulas['Matriz_Consecuencias']['A'+str(row)].value,'status':'numeric_text_rule' if row in [2,3,4,5,9] else 'qualitative_text_rule','condition_cell':'C'+str(row),'condition':formulas['Matriz_Consecuencias']['C'+str(row)].value,'effect_cell':'D'+str(row),'effect':formulas['Matriz_Consecuencias']['D'+str(row)].value})
    normalized=[{'sheet':s,'normalized_to_CV1000':norm,'representative':addresses[0],'cells':addresses,'count':len(addresses)} for (s,norm),addresses in families.items()]
    excel_errors={'#DIV/0!','#VALUE!','#REF!','#NAME?','#NUM!','#N/A','#NULL!','#SPILL!','#CALC!'}
    # This workbook stores four formula-error caches as strings (t="str"),
    # so data_type=='e' alone incorrectly reports zero cached error results.
    errors=[{**c,'cache_storage_type':cache[c['sheet']][c['cell']].data_type} for c in all_cells if cache[c['sheet']][c['cell']].data_type=='e' or (isinstance(c['cached'],str) and c['cached'] in excel_errors)]
    report={'source':str(SOURCE.relative_to(ROOT)),'sha256_before':before,'sha256_after':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'mode':'read-only dual load; original cache is not a fresh recalculation','counts_by_sheet':counts,'formula_count':sum(x.get('formula',0) for x in counts),'formula_family_count':len(families),'family_normalization':'openpyxl Translator moves copied formulas to CV1000; absolute references and constants preserved; families are structural, not proof of mathematical correctness.','cached_errors':errors,'comparisons':comparisons,'comparison_count':len(comparisons),'differences':[c for c in comparisons if not c['equal']],'formula_families':normalized,'consequences':consequence_rows,'original_ambiguities':contradictions,'rules_absent':absent,'all_nonempty_cells':all_cells}
    assert report['sha256_after']==before
    dest=ROOT/'reports/auditoria_excel_original_20261023.json'
    dest.write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
    lines=['# Auditoría reproducible del Excel original — 23/10/2026','',f'Fuente intacta: `{report["source"]}`. SHA-256 antes/después: `{before}`. Skill xlsx aplicada; dos cargas independientes openpyxl. No se ejecutó COM ni se guardó el libro.','',f'Se inventariaron {len(all_cells)} celdas no vacías, {report["formula_count"]} fórmulas y {len(families)} familias estructurales. Todas las fórmulas/cachés y pertenencias a familias están en el JSON acompañante. Errores en caché original: {len(errors)}.','',f'Comparaciones nuevas contra SIDE: {len(comparisons)}; diferencias: {len(report["differences"])}. Incluyen productividades, consumos, conversiones, precios, todos los componentes laborales de ambos regímenes, todos los activos de ambos regímenes, parámetros de los tres mercados y catálogo completo. Son comparaciones con caché original, no prueba de recálculo ni cobertura integral de transiciones de caja.','','## Cobertura por hoja','','| Hoja | Fórmulas | Literales numéricos | Textos |','|---|---:|---:|---:|']
    for c in counts: lines.append(f'| {c["sheet"]} | {c.get("formula",0)} | {c.get("numeric_literal",0)} | {c.get("text",0)} |')
    lines+=['','Errores originales guardados como texto en la caché (tipo `s`, no `e`): '+', '.join(f'`{c["sheet"]}!{c["cell"]}={c["cached"]}`' for c in errors)+'. La detección comprueba tanto el tipo como los códigos textuales de error.','']
    lines+=['','## Diferencias observadas contra SIDE','']
    for d in report['differences']: lines.append(f'- {d["label"]}: `{d["excel_sheet"]}!{d["excel_cell"]}` original `{d["excel_cached"]}`, SIDE `{d["side"]}`; fuente `{d["side_source"]}`.')
    if not report['differences']: lines.append('Sin diferencias en las comparaciones enumeradas. Esto no convierte reglas ausentes en fórmulas originales.')
    lines+=['','## Consecuencias: fórmula versus texto','','Ninguna celda de Matriz_Consecuencias es una fórmula. Los efectos numéricos son reglas escritas en texto.','']
    for c in consequence_rows: lines.append(f'- **{c["id"]}** ({c["status"]}): C{int(c["condition_cell"][1:])} `{c["condition"]}`; {c["effect_cell"]} `{c["effect"]}`.')
    lines+=['','CONS-03 usa lo **planeado**, no la capacidad ni la producción ya limitada. CONS-06 evalúa **cantidad producida** mayor que capacidad; cambiarlo por cantidad deseada cambia el disparador textual. No especifica tarifa de horas extra. No hay redondeo en esas reglas textuales.','','## Ambigüedades y contradicciones originales','']
    for c in contradictions: lines.append(f'- **{c["id"]}** ({c["status"]}): {c["detail"]} Celdas: '+', '.join(f'`{s}!{a}`' for s,a in c['cells'])+'.')
    lines+=['','## Reglas que el original no determina','']
    for c in absent: lines.append(f'- **{c["rule"]}**: {c["detail"]}')
    lines+=['','Las seis diferencias pequeñas de nómina NO_MYPE (si siguen presentes) reflejan cachés originales guardadas a seis decimales. Las fórmulas no contienen ROUND: no se debe redondear SIDE para copiar una caché truncada. El recálculo COM debe resolverlas.','','## Reproducción y límites','','Ejecutar `tests/audit_excel_original_20261023.py` con Python que tenga openpyxl, desde SIDE1. Usa Node para leer configuración y ejecutar costoLaboral/costoMaquinaria actuales. JSON siempre UTF-8 explícito. El proceso verifica hash antes/después y no modifica el Excel.','', 'La cobertura de familias es un inventario exhaustivo, no una ejecución independiente de cada fórmula. Las diferencias de fórmulas entre hojas y las reglas ausentes requieren criterio explícito del integrador; no se corrige el original silenciosamente. El oráculo fresco COM y las simulaciones quedan a cargo del proceso de integración para mantener un único motor Excel activo.']
    (ROOT/'reports/auditoria_excel_original_20261023.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    print(json.dumps({k:report[k] for k in ['formula_count','formula_family_count','comparison_count','differences']},ensure_ascii=True,indent=2))

def build_com_plans(book,quantity,costs):
    """Prepare original-only perturbations; execution is owned by the integrator."""
    def val(sheet,cell,mutations): return mutations.get(sheet,{}).get(cell,book[sheet][cell].value)
    def records(mutations):
        checks=[]
        personal={area:[[val(quantity,openpyxl.utils.get_column_letter(col)+str(row),mutations) or 0 for col in range(6,9)] for row in rows] for area,rows in [('corte',range(5,8)),('ensamble',range(11,14)),('acabado',range(17,20))]}
        inp={'personal':personal,'jefes':val(quantity,'F21',mutations)}
        for cell,path in [('I8','capacidadPorArea.corte'),('I14','capacidadPorArea.ensamble'),('I20','capacidadPorArea.acabado'),('K5','capacidadDiaria'),('L5','capacidadCicloBase'),('M5','factorEficienciaProduccion'),('N5','produccionCapacidad')]:
            checks.append({'key':'quantity__'+cell,'module':'produccion','function':'calcularCapacidad','input':inp,'output_path':path,'scope':'original_quantity_formula'})
        cost_personal={area:[[val(costs,openpyxl.utils.get_column_letter(col)+str(row),mutations) or 0 for col in range(6,9)] for row in rows] for area,rows in [('corte',range(5,8)),('ensamble',range(11,14)),('acabado',range(17,20))]}
        cost_productivity={area:[val(costs,'E'+str(row),mutations) for row in rows] for area,rows in [('corte',range(5,8)),('ensamble',range(11,14)),('acabado',range(17,20))]}
        for cell,path in [('I8','capacidadPorArea.corte'),('I14','capacidadPorArea.ensamble'),('I20','capacidadPorArea.acabado'),('K5','capacidadDiaria'),('L5','capacidadCicloBase'),('M5','factorEficienciaProduccion'),('N5','produccionCapacidad')]:
            checks.append({'key':'costs__'+cell,'module':'produccion','function':'calcularCapacidad','input':{'personal':cost_personal,'jefes':val(costs,'F21',mutations)},'overrides':{'DIAS_CICLO':30,'PRODUCTIVIDAD':cost_productivity},'output_path':path,'scope':'original_cost_sheet_capacity_literal'})
        target_lots=[{'cantidad':val(quantity,'D'+str(37+i),mutations),'molde':code} for i,code in enumerate(['INF-08','INF-09','INF-10'])]
        for i,cells in enumerate([['D42','G42','J42'],['E43','H43','K43'],['F44','I44','L44']]):
            for family,cell in zip(['cuero','accesorios','hilo'],cells):
                checks.append({'key':'quantity__'+cell,'module':'produccion','function':'calcularProduccion','input':{'lotes':target_lots},'output_path':f'lotes.{i}.requerimiento.{family}','scope':'original_quantity_requirements_without_stock'})
        score_input={'seleccion':[val('Seleccion_Ciclo0','B'+str(row),mutations) for row in range(2,16)],'lotes':[]}
        for row in range(2,5):
            score_input['lotes'].append({key:val('Lotes_Produccion',col+str(row),mutations) for key,col in [('segmento','B'),('cantidad','C'),('cuero','D'),('accesorios','E'),('hilo','F'),('molde','G')]})
            for col,field in [('H','costoUnitario'),('I','calidad'),('J','umbral'),('K','resultado'),('L','mensaje')]:
                checks.append({'key':'lots__'+col+str(row),'module':'puntajes','function':'calcularPuntajes','input':score_input,'output_path':f'lotes.{row-2}.{field}','scope':'original_lot_formula','blank_string_is_null':True})
        for cell,path in [('B6','calidadEmpresaMotor'),('D6','niveles.calidad'),('B9','puntajeEficiencia'),('D9','niveles.eficiencia'),('B10','reputacionMatriz'),('D10','niveles.reputacion'),('B11','saludFinanciera'),('D11','niveles.salud')]:
            checks.append({'key':'scores__'+cell,'module':'puntajes','function':'calcularPuntajes','input':score_input,'output_path':path,'scope':'original_score_formula','blank_string_is_null':True})
        for cell,path in [('C5','totalCantidades'),('H5','costoInsumos')]:
            checks.append({'key':'lots__'+cell,'module':'puntajes','function':'calcularPuntajes','input':score_input,'output_path':path,'scope':'original_lot_total_formula'})
        for regime,cols in [('MYPE',['U','Y','AC']),('NO_MYPE',['V','Z','AD'])]:
            for role,col in list(zip(['corte','ensamble','acabado'],cols))+[('jefatura',cols[0]),('limpieza',cols[1])]:
                profiles=[(1,5),(2,14),(3,23)] if role in ['corte','ensamble','acabado'] else [(1,33)]
                for level,start in profiles:
                    salaries={r:[val(costs,c+str(row),mutations) for row in [5,14,23]] for r,c in [('corte','U'),('ensamble','Y'),('acabado','AC')]}
                    salaries.update({'jefatura':val(costs,'U33',mutations),'limpieza':val(costs,'Y33',mutations)})
                    for field,offset in [('base',0),('essalud',1),('gratificacion',2),('cts',3),('vacaciones',4),('sumaAnual',5),('pagoMensual',6)]:
                        checks.append({'key':'costs__'+col+str(start+offset),'module':'costos','function':'costoLaboral','arguments':[role,level,regime],'overrides':{'SUELDOS':salaries},'output_path':field,'scope':'original_payroll_formula'})
        for code,col,row,mold in [('mesa','U',45,False),('INF-02','Y',45,False),('INF-03','Y',53,False),('INF-04','Y',61,False),('INF-05','AC',45,False),('INF-06','AC',53,False),('INF-07','AC',61,False),('INF-08','U',53,True),('INF-09','U',60,True),('INF-10','U',67,True)]:
            for regime in ['MYPE','NO_MYPE']:
                cl=col if regime=='MYPE' else openpyxl.utils.get_column_letter(openpyxl.utils.column_index_from_string(col)+1)
                offsets=[('capacitacion',1),('mantenimientoAnual',2),('depreciacionAnual',3),('sumaAnual',4)] if mold else [('instalacion',1),('capacitacion',2),('mantenimientoAnual',3),('depreciacionAnual',4),('sumaAnual',5)]
                dep_offset=3 if mold else 4
                dep=book[costs][cl+str(row+dep_offset)].value
                life=int(re.search(r'/([0-9]+)',dep).group(1))
                overrides={'MAQUINARIA':{code:{'costo':val(costs,cl+str(row),mutations),'vida':life,'molde':mold}}}
                for field,offset in offsets: checks.append({'key':'costs__'+cl+str(row+offset),'module':'costos','function':'costoMaquinaria','arguments':[code,regime,{},overrides],'output_path':field,'scope':'original_asset_formula'})
        # Original model2 quality is a reputation proxy, independent of any
        # external quality input. Literal override is explicit, never implicit.
        econ_sum=sum(val(quantity,'D'+str(row),mutations) for row in range(128,132))
        for segment,start,pricecol,repcol in [('SEG-01',128,'C','D'),('SEG-02',141,'C','D'),('SEG-03',154,'C','D')]:
            companies=[{'id':str(i),'precio':val(quantity,pricecol+str(start+i),mutations),'reputacionMercado':val(quantity,repcol+str(start+i),mutations),'calidadEmpresaMotor':0} for i in range(4)]
            for i in range(4):
                for col,path in [('E','calidadMercado'),('F','puntajePrecio'),('H','puntajeTotal'),('I','porcentaje'),('J','asignado')]:
                    checks.append({'key':'quantity__'+col+str(start+i),'module':'mercado','function':'repartoMercado','arguments':[companies,segment,{'CALIDAD_MERCADO_MODO':'excel_literal'},{'sumaReputacionEconomico':econ_sum}],'output_path':str(i)+'.'+path,'scope':'original_market2_literal','expected_domain_error':segment=='SEG-01' and len({x['precio'] for x in companies})==1})
        return checks
    variants=[
        ('original_baseline',{}),
        ('market_distinct_reputations',{quantity:{'D128':1,'D129':2,'D130':3,'D131':4,'D141':8,'D142':7,'D143':6,'D144':5,'D154':9,'D155':6,'D156':3,'D157':1},costs:{'D118':1,'D119':2,'D120':3,'D121':4,'D131':8,'D132':7,'D133':6,'D134':5,'D144':9,'D145':6,'D146':3,'D147':1}}),
        ('market_economic_price_tie',{quantity:{'C128':40,'C129':40,'C130':40,'C131':40},costs:{'C118':40,'C119':40,'C120':40,'C121':40}}),
        ('market_negative_ideal_price_score',{quantity:{'C141':500,'C154':1000},costs:{'C131':500,'C144':1000}}),
        ('staff_none',{quantity:{**{col+str(row):0 for row in [5,6,7,11,12,13,17,18,19] for col in ['F','G','H']},'F21':0}}),
        ('staff_n3_industrial',{quantity:{**{col+str(row):0 for row in [5,6,7,11,12,13,17,18,19] for col in ['F','G','H']},'F7':2,'H13':2,'H19':2,'F21':1}}),
        ('staff_high_efficiency_unclamped',{quantity:{'F7':20,'F13':20,'F19':20,'F21':2}}),
        ('salary_perturbed',{costs:{'U5':1837,'V5':1837,'U14':2613,'V14':2613,'Y14':3127,'Z14':3127,'AC23':4871,'AD23':4871}}),
        ('assets_perturbed',{costs:{'U45':3100,'V45':3100,'Y53':8123,'Z53':8123,'AC61':10123,'AD61':10123,'U67':3901,'V67':3901}}),
        ('bom_fractional',{quantity:{'D37':1,'D38':2,'D39':3},costs:{'D46':1,'D47':2,'D48':3}}),
        ('market_model1_distinct_populations',{quantity:{'B89':10000,'B90':0.5,'B103':3000,'B104':0.3},costs:{'B79':10000,'B80':0.5,'B93':3000,'B94':0.3}}),
        ('quality_lot_proxy_independence',{'Lotes_Produccion':{'D2':'INS-03','E2':'INS-05','F2':'INS-08','G2':'INF-10','C2':120}})
    ]
    outputs={}
    for sheet,prefix in [(quantity,'quantity'),(costs,'costs'),('Lotes_Produccion','lots'),('Motor_Puntajes','scores')]:
        for row in book[sheet]:
            for cell in row:
                if cell.data_type=='f': outputs[prefix+'__'+cell.coordinate]=[sheet,cell.coordinate]
    plans=[]
    for name,mutations in variants:
        checks=records(mutations)
        case_outputs=dict(outputs)
        source_sheets={'quantity':quantity,'costs':costs,'lots':'Lotes_Produccion','scores':'Motor_Puntajes'}
        for check in checks:
            prefix,cell=check['key'].split('__')
            case_outputs.setdefault(check['key'],[source_sheets[prefix],cell])
        plans.append({'nombre':name,'entradas':mutations,'salidas':case_outputs,'side_checks':checks,'note':'Copy untouched original, apply only inputs, CalculateFullRebuild, preserve Excel errors as errors. No added sheets. Includes exact formula snapshots separately in audit report.'})
    arithmetic=[]
    for i,(initial,bought,used) in enumerate([(0.1,0.7,0.8),(0.1,0.2,0.3),(0.3,0.6,0.9),(0.6,0.3,0.9)],1):
        arithmetic.append({'nombre':'operand_injection_inventory_'+str(i),'entradas':{costs:{'D54':initial,'E54':bought,'F54':used}},'salidas':{'costs__G54':[costs,'G54']},'formula_modified_in_copy':True,'original_formula_overwritten':{'sheet':costs,'cell':'F54','formula':book[costs]['F54'].value},'expected_arithmetic_identity':0,'note':'Arithmetic unit experiment only. F54 originally is a formula and is intentionally replaced with its operand result in the output COPY. This is not an original-input scenario and must not be counted as unmodified-source oracle coverage. G54 remains the original formula D54+E54-F54. Original workbook untouched.'})
    path=ROOT/'tests/oracle/auditoria_original_20261023_planes.json'
    path.write_text(json.dumps({'source':'.codex-work/DECISIONES_SIDE_16.xlsx','sha256':SHA,'planes':plans,'planes_operando_inyectado':arithmetic},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

if __name__=='__main__': main()

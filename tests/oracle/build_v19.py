"""Build corrected v19 from the explicitly authorized v16; never overwrite source.

Run with bundled Python from SIDE1. Recalculation is a separate, mandatory step.
All oracle expected numbers come from workbook formulas after Excel/LO recalc.
"""
from pathlib import Path
import json
import re
import hashlib
from copy import copy
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / '.codex-work/DECISIONES_SIDE_16.xlsx'
DEST = ROOT / 'docs/DECISIONES_SIDE_19.xlsx'
book = openpyxl.load_workbook(SOURCE)
changes = []


def put(sheet, cell, value, reason):
    ws = book[sheet] if isinstance(sheet, str) else sheet
    before = getattr(ws[cell].value, 'text', ws[cell].value)
    if before == value:
        return
    changes.append([ws.title, cell, before, value, reason])
    ws[cell] = value


PROD = 'CANTIDAD DE PRODUCCIÓN'
COST = 'Costos De Producción'
for ws in book:
    for row in ws:
        for cell in row:
            formula = getattr(cell.value, 'text', cell.value)
            if cell.data_type == 'f' and isinstance(formula, str) and '$57' in formula:
                put(ws, cell.coordinate, formula.replace('$57', '$69'), 'E1: incluir las 68 opciones')

cat = book['Catalogo_Decisiones']
for row in range(2, 70):
    for col in range(7, 12):
        c = cat.cell(row, col)
        if isinstance(c.value, str) and re.fullmatch(r'=\+\d+', c.value):
            put(cat, c.coordinate, int(c.value[2:]), 'E2: fórmula constante a literal numérico; original sí evaluaba como número')
        elif isinstance(c.value, str) and c.value.startswith('S/'):
            put(cat, c.coordinate, float(c.value[2:].replace(',', '')), 'E2: moneda textual a número, formato monetario')
            c.number_format = '"S/"#,##0.00'

selection = book['Seleccion_Ciclo0']
for row, label in [(24, 'Método de depreciación'), (25, 'Garantía'), (26, 'Mantenimiento'), (27, 'Política de inventarios')]:
    put(selection, f'A{row}', label, 'E1: exponer grupo existente del catálogo, sin decisión nueva')
    for col, catcol in [('C','D'),('D','F'),('E','G'),('F','H'),('G','I'),('H','J'),('I','K')]:
        fallback = '""' if col in 'CD' else '0'
        put(selection, f'{col}{row}', f'=IFERROR(INDEX(Catalogo_Decisiones!${catcol}$2:${catcol}$69,MATCH($B{row},Catalogo_Decisiones!$A$2:$A$69,0)),{fallback})', 'E1: búsqueda de grupos existentes')
for dv in selection.data_validations.dataValidation:
    if dv.formula1 and '$57' in dv.formula1:
        dv.formula1 = dv.formula1.replace('$57', '$69')
selection.add_data_validation(DataValidation(type='list',formula1='Catalogo_Decisiones!$A$2:$A$69'))
selection.data_validations.dataValidation[-1].add('B24:B27')

lots = book['Lotes_Produccion']
for row in range(2, 5):
    original_quality = getattr(lots[f'I{row}'].value, 'text', lots[f'I{row}'].value)
    put(lots, f'I{row}', f'=MAX(0,MIN(10,{original_quality[1:]}))', 'E9: calidad acotada 0–10')
    put(lots, f'J{row}', f'=IF(B{row}="SEG-03",7,IF(B{row}="SEG-02",4,0))', 'Compatibilidad: IF anidado en vez de IFS sin prefijo')
    put(lots, f'M{row}', [42.5,160,250][row-2], 'E3: precio elegido; supuesto inicial por segmento, editable')
    put(lots, f'N{row}', f'=IF(B{row}="","",IF(AND(M{row}>=IF(B{row}="SEG-01",40,IF(B{row}="SEG-02",120,220)),M{row}<=IF(B{row}="SEG-01",45,400)),"Válido","Fuera de rango"))', 'E3: validación visible de precio')
    # Standard upper range is 200, premium 400.
    put(lots, f'N{row}', f'=IF(B{row}="","",IF(AND(M{row}>=IF(B{row}="SEG-01",40,IF(B{row}="SEG-02",120,220)),M{row}<=IF(B{row}="SEG-01",45,IF(B{row}="SEG-02",200,400))),"Válido","Fuera de rango"))', 'E3: límites económicos40–45,medio120–200,premium220–400')
    put(lots, f'O{row}', f'=M{row}*IF(B{row}="SEG-03",IF(I{row}<7,0.65,1.4),1)', 'E9: Premium<7 penaliza, >=7 premia')
put(lots, 'M1', 'Precio elegido (S/ por unidad)', 'E3: único campo nuevo permitido')
put(lots, 'N1', 'Validación de precio', 'E3')
put(lots, 'O1', 'Precio efectivo tras consecuencias', 'E9/CONS01/02')
dv = DataValidation(type='custom',formula1='AND(M2>=IF(B2="SEG-01",40,IF(B2="SEG-02",120,220)),M2<=IF(B2="SEG-01",45,IF(B2="SEG-02",200,400)))')
dv.errorTitle='Precio fuera de rango';dv.error='Económico40–45; Medio120–200; Premium220–400';dv.showErrorMessage=True;dv.errorStyle='stop'
lots.add_data_validation(dv);dv.add('M2:M4')

for cell, col in [('B9','G'),('B10','H'),('B11','I')]:
    put('Motor_Puntajes',cell,f'=MAX(0,MIN(10,SUM(Seleccion_Ciclo0!{col}2:{col}15,Seleccion_Ciclo0!{col}24:{col}27)))','E6/E9: catálogo única fuente, tope0–10')
put('Matriz_Consecuencias','C2','< 7 y Segmento = Premium','E9: cerrar hueco4–6.99; supuesto documentado')
for sn, rep_rows in [(PROD,[75,123]),(COST,[65,113])]:
    for row in rep_rows:
        put(sn,f'A{row}','Reputación canónica del catálogo (0–10)','E6: retirar mezcla producción/publicidad/puntajes')
        for col in 'CDEF':
            put(sn,f'{col}{row}',0,'E6: componente histórico obsoleto, sin doble publicidad')
        put(sn,f'G{row}','=Motor_Puntajes!B10','E6: única reputación de decisiones, acotada')
    put(sn,'M10','=IFERROR(L10/L9,0)','Proteger ausencia de personal')
    put(sn,'M5','=MIN(1,K7+SUM(F7,F13:H13,F19:H19)/100+4%*F21)','E7: bono por TODOS N3; SUPUESTO eficiencia máxima100%')
for row in [5,6,7,11,12,13,17,18,19]:
    put(COST,f'E{row}',f"='{PROD}'!E{row}",'E7: productividad canónica CANTIDAD')
put(COST,'L5','=K5*24','E7:24 días/ciclo')
for row in range(37,40):
    put(PROD,f'E{row}',f'=IFERROR(ROUNDDOWN(D{row}*E35/D36,0),0)','Proteger plan total cero')
for row in range(56,59):
    put(PROD,f'K{row}',f'=Lotes_Produccion!I{row-54}','E8: recuperar calidad desde puntaje lote, no fechas')
    book[PROD][f'K{row}'].number_format='0.00'

# Replace both duplicate v2 examples with the same canonical equations.
for sn, offset in [(PROD,0),(COST,-10)]:
    ws=book[sn]
    for group, start in enumerate([128+offset,141+offset,154+offset]):
        end=start+3; total=end+1; real=start+8
        marketbase=start+6; penetration=start+7
        for r in range(start,end+1):
            lotrow=2+min(r-start,2)
            put(ws,f'N{r}',f'=Lotes_Produccion!I{lotrow}','E5a: calidad real lote0–10, editable ejemplo por empresa')
            put(ws,f'O{r}',f'=Lotes_Produccion!C{lotrow}','E5e: stock inicial+producción disponible, ejemplo')
            put(ws,f'P{r}',f'=Lotes_Produccion!C{lotrow}','E5f: producción del segmento; 0 excluye')
            put(ws,f'E{r}',f'=MAX(0,MIN(10,N{r}))*10','E5a/b: calidad normalizada0–100, eliminar división reputaciones')
            if group==0:
                score=f'IF(MAX($C${start}:$C${end})=MIN($C${start}:$C${end}),100,100*(MAX($C${start}:$C${end})-C{r})/(MAX($C${start}:$C${end})-MIN($C${start}:$C${end})))'
            else:
                ideal=100 if group==1 else 250
                score=f'100-ABS(C{r}-{ideal})/{ideal}*100'
            put(ws,f'F{r}',f'=MAX(0,MIN(100,{score}))','E5d: límites y empate')
            put(ws,f'G{r}',f'=MAX(0,MIN(10,D{r}))*10','E9: reputación0–100')
            wr=start-4
            put(ws,f'H{r}',f'=IF(P{r}>0,$M${wr}*F{r}+$M${wr+1}*G{r}+$M${wr+2}*E{r},0)','E5f: solo compiten productores')
            put(ws,f'I{r}',f'=IFERROR(H{r}/SUM($H${start}:$H${end}),0)','E5b: denominador del propio segmento')
            put(ws,f'J{r}',f'=ROUND(I{r}*$B${real},0)','E5: demanda asignada v2')
            put(ws,f'Q{r}',f'=MIN(MAX(0,O{r}),J{r})','E5e: venta limitada a stock; demanda residual perdida')
            put(ws,f'S{r}',f'=C{r}*'+(f'IF(N{r}<7,0.65,1.4)' if group==2 else '1'),'E9: precio efectivo Premium por calidad')
            put(ws,f'R{r}',f'=Q{r}*S{r}','E3: ingreso con precio efectivo')
        for col,label in [('N','Calidad lote0–10'),('O','Stock disponible'),('P','Producción segmento'),('Q','Venta MIN stock'),('R','Ingresos S/'),('S','Precio efectivo S/')]:
            put(ws,f'{col}{start-1}',label,'E5: entradas/salidas explícitas')
        put(ws,f'B{real}',f'=B{marketbase}*B{penetration}','E5c: mercado real calculado')
    put(ws,'A170','V2 CANÓNICA. Bloques v1 históricos obsoletos. Demanda no atendida se pierde; solo productores compiten.','E4/E5')

assumptions=book.create_sheet('Supuestos')
assumptions.append(['ID','Parámetro','Valor / regla','Estado / fundamento'])
rows=[
('FUENTE','Base autorizada','v16; v18 no disponible','Autorizado por usuario; hash en control'),
('E3','Precio por segmento','40–45 /120–200 /220–400; ideales100/250','Rangos explícitos usuario; defaults42.5/160/250 supuesto'),
('E4','Mercado','v2; pesos .7/.2/.1,.4/.35/.25,.1/.5/.4;3500/3500/800','Canónico por usuario'),
('E5','Empate económico','PrecioScore100 todos; scores acotados0–100','Regla solicitada'),
('E5e','Demanda no atendida','Se pierde, sin reasignación','Regla solicitada'),
('E5f','Competidores','Solo empresas que produjeron lote en segmento','Regla solicitada; inventario antiguo solo vende si también produce'),
('E6','Reputación','Catálogo una opción por grupo; clamp0–10, luego CONS01/02 clamp','SUPUESTO para aprobación; publicidad no se cuenta otra vez'),
('E7','Capacidad','24d,90%base,+1pp por N3,+4pp jefatura; corte5/7/9,ensamble6/8/12,acabado6/7/10','Canon CANTIDAD solicitado'),
('E9','Calidad y puntajes','Tope0–10; Premium<7 CONS01 y>=7 CONS02','SUPUESTO hueco Premium resuelto'),
('E9','Personal mixto','Aporte máximo por grupo, no por cada empleado','SUPUESTO para evitar inflar puntajes por volumen'),
('E11','Canal','Cobertura base1 para física/digital/mixto','SUPUESTO configurable neutral, pendiente aprobación'),
('E11','Tiendas','+5% demanda por tienda adicional, máximo+25%','SUPUESTO configurable, pendiente aprobación'),
('E11','Marketing','A1,B1.12,C1.20','B+12% del catálogo; A/C SUPUESTO; no publicidad extra'),
('E11','Crédito','Inicio20%,+5pp/ciclo,+5pp analista,máximo70%; elección explícita prevalece; cobranza ciclo siguiente','SUPUESTO heredado configurable, sin mora'),
('E11','Garantía','0/5/10 soles por unidad vendida','Costos catálogo'),
('E11','Capacitación','A0%,B8%,C12%capacidad','B+8% catálogo; otras SUPUESTO'),
('E11','Mantenimiento','0/1500/3000 soles del catálogo; no mejora capacidad automáticamente','Factores de capacidad0%; SUPUESTO neutral'),
('E11','Inventarios','Conservar físicos; política no crea ni borra stock','SUPUESTO; costos catálogo0/1000/2500'),
('E11','Depreciación','Solo contabilidad; vida productiva100000u; acelerada2×lineal','Regla solicitada; vida y aceleración SUPUESTOS'),
('E11','Comisión','1% de ingresos efectivos','SUPUESTO heredado configurable'),
('E7','Eficiencia operativa','Máximo100%, contando todos los trabajadores N3','SUPUESTO coherente con runtime'),
('E11','Materiales mixtos','FIFO de opciones; calidad media ponderada de lo realmente consumido','SUPUESTO para mezclas no especificadas por lote en UI'),
('CONS05','Riesgo bancario','+3 puntos porcentuales al siguiente préstamo','SUPUESTO configurable'),
('CONS06','Sobrecapacidad','Sobrecosto25%planilla×fracción exceso, máximo100%; retraso proporcional','SUPUESTO configurable'),
('CONS07','Almacenaje','S/2 por unidad final sin vender','SUPUESTO configurable'),
('REDONDEO','Unidades','ROUND demanda, floor producción; MIN stock; ingresos2decimales','SUPUESTO para consistencia contable'),
]
for r in rows:assumptions.append(r)

# Independent formula oracle: 9 cases, four companies per segment, no JS output imported.
oracle=book.create_sheet('Oraculo_v19')
headers=['caso','empresa','segmento','precio','reputacion','calidad0a10','stock','producido','bajaPrevios','pesoP','pesoR','pesoC','mercado','ideal','precioScore','puntaje','participacion','demanda','vendidas','factorPrecio','precioEfectivo','ingresos','demandaPerdida','reputacionFinal']
oracle.append(headers)
cases=[('normal_economico','SEG-01'),('normal_medio','SEG-02'),('normal_premium','SEG-03'),('empate','SEG-01'),('sin_stock','SEG-02'),('no_produjo','SEG-02'),('premium_incoherente','SEG-03'),('demanda_no_atendida','SEG-01'),('reputacion_baja_2c','SEG-02')]
for name,segment in cases:
    begin=oracle.max_row+1;end=begin+3
    params={'SEG-01':([40,42,43,45],(.7,.2,.1),3500,0),'SEG-02':([120,140,160,200],(.4,.35,.25),3500,100),'SEG-03':([220,250,300,400],(.1,.5,.4),800,250)}[segment]
    for i in range(4):
        r=oracle.max_row+1;prices,weights,market,ideal=params
        price=42 if name=='empate' else prices[i]
        rep=2 if name=='reputacion_baja_2c' and i==0 else [7,8,9,6][i]
        quality=5 if name=='premium_incoherente' and i==0 else [8,7,9,6][i]
        stock=0 if name=='sin_stock' and i==0 else (10 if name=='demanda_no_atendida' else 1000)
        produced=0 if name=='no_produjo' and i==0 else 100
        oracle.append([name,chr(65+i),segment,price,rep,quality,stock,produced,1 if name=='reputacion_baja_2c' else 0,*weights,market,ideal])
        maxp='MAX('+','.join(f'IF($H${k}>0,$D${k},-1E99)' for k in range(begin,end+1))+')'
        minp='MIN('+','.join(f'IF($H${k}>0,$D${k},1E99)' for k in range(begin,end+1))+')'
        # Eligible-price extrema exclude firms with no produced lot.
        oracle.cell(r,15,f'=IF(H{r}<=0,0,MAX(0,MIN(100,IF(C{r}="SEG-01",IF({maxp}={minp},100,100*({maxp}-D{r})/({maxp}-{minp})),100-ABS(D{r}-N{r})/IF(N{r}=0,1,N{r})*100))))')
        oracle.cell(r,20,f'=IF(C{r}="SEG-03",IF(F{r}<7,0.65,1.4),1)')
        oracle.cell(r,24,f'=MAX(0,MIN(10,E{r}+IF(C{r}="SEG-03",IF(F{r}<7,-2,2),0)))')
        oracle.cell(r,16,f'=IF(H{r}>0,J{r}*O{r}+K{r}*X{r}*10+L{r}*MAX(0,MIN(10,F{r}))*10,0)')
        oracle.cell(r,17,f'=IFERROR(P{r}/SUM($P${begin}:$P${end}),0)')
        oracle.cell(r,18,f'=ROUNDDOWN(ROUND(Q{r}*M{r},0)*IF(AND(X{r}<3,I{r}>=1),0.8,1),0)')
        oracle.cell(r,19,f'=MIN(G{r},R{r})')
        oracle.cell(r,21,f'=D{r}*T{r}')
        oracle.cell(r,22,f'=ROUND(S{r}*U{r},2)')
        oracle.cell(r,23,f'=MAX(0,R{r}-S{r})')

components=book.create_sheet('Componentes_v19')
components.append(['bloque','nombre','valor_formula','fuente'])
for area, cells in [('corte',['E5','E6','E7']),('ensamble',['E11','E12','E13']),('acabado',['E17','E18','E19'])]:
    for level,c in enumerate(cells,1):components.append(['capacidad',f'{area}_nivel{level}',f"='{PROD}'!{c}",f'{PROD}!{c}'])
for name,c in [('eficiencia','M5'),('capacidad','N5'),('produccion_lote1','E56'),('produccion_lote2','E57'),('produccion_lote3','E58')]:components.append(['produccion',name,f"='{PROD}'!{c}",f'{PROD}!{c}'])
for name,c in [('calidad1','B3'),('calidad2','B4'),('calidad3','B5'),('calidadEmpresa','B6'),('eficienciaPuntaje','B9'),('reputacion','B10'),('saludFinanciera','B11')]:components.append(['puntajes',name,f'=Motor_Puntajes!{c}',f'Motor_Puntajes!{c}'])
for level in range(3):
    rows=[5+level,11+level,17+level]
    formula=f'=ROUND(MIN(4*\'{PROD}\'!E{rows[0]},4*\'{PROD}\'!E{rows[1]}*1.5,4*\'{PROD}\'!E{rows[2]}*1.5)*24*MIN(1,0.9+{12 if level==2 else 0}*0.01+0.04),0)'
    components.append(['capacidad',f'nivel{level+1}_capacidad',formula,'4operarios/área,4máquinas,ensamble/acabado industriales,jefatura'])

full=book.create_sheet('Cierre_oraculo_v19')
full.append(['empresa','segmento','objetivo','capacidad','producido','calidad','reputacion','precio','precioScore','puntaje','participacion','asignado','vendido','precioEfectivo','ingresos','consumoCuero','consumoAccesorios','consumoHilo','ingresoEmpresa','cuentasCobrar','ingresoCaja','almacenaje','eficiencia','saludFinanciera'])
full['AA1']='Códigos de entradas del fixture; no resultados JS'
codes=['INF-04','INF-07','INF-09','RH-03','RH-06','RH-09','RH-10A','COM-01B','FIN-01B']
for i,code in enumerate(codes,2):full.cell(i,27,code)
for col,catcol in [(28,'I'),(29,'J'),(30,'K')]:
    full.cell(2,col,'=MAX(0,MIN(10,'+'+'.join(f'SUMIF(Catalogo_Decisiones!$A$2:$A$69,$AA${r},Catalogo_Decisiones!${catcol}$2:${catcol}$69)' for r in range(2,11))+'))')
lookup=lambda code,col:f'SUMIF(Catalogo_Decisiones!$A$2:$A$69,"{code}",Catalogo_Decisiones!${col}$2:${col}$69)'
for company in range(4):
    first=2+company*3;last=first+2
    for i,segment in enumerate(['SEG-01','SEG-02','SEG-03']):
        r=first+i;target=30 if company==3 else (90 if i==company else 0)
        for col,value in enumerate([str(company),segment,target],1):full.cell(r,col,value)
        full.cell(r,4,f'=ROUND(MIN(4*\'{PROD}\'!E7,4*\'{PROD}\'!E13*1.5,4*\'{PROD}\'!E19*1.5)*24*MIN(1,0.9+12*0.01+0.04),0)')
        # Materials are abundant for these real four-company fixtures; physical bound remains explicit.
        full.cell(r,5,f'=MIN(C{r},D{r},ROUNDDOWN(250/'+str([.3,.4,.7][i])+',0),1000,ROUNDDOWN(10000/'+str([6,7,13][i])+',0))')
        full.cell(r,6,'=MAX(0,MIN(10,'+'+'.join(lookup(code,'H') for code in ['INS-03','INS-05','INS-08',['INF-08','INF-09','INF-10'][i]])+'))')
        full.cell(r,7,f'=MAX(0,MIN(10,$AC$2+IF($C${last}>0,2,0)))')
        full.cell(r,8,[42.5,160,250][i])
        full.cell(r,9,100 if i==0 else f'=MAX(0,MIN(100,100-ABS(H{r}-{100 if i==1 else 250})/{100 if i==1 else 250}*100))')
        weights=[(.7,.2,.1),(.4,.35,.25),(.1,.5,.4)][i]
        full.cell(r,10,f'=IF(E{r}>0,{weights[0]}*I{r}+{weights[1]}*G{r}*10+{weights[2]}*F{r}*10,0)')
        full.cell(r,11,f'=IFERROR(J{r}/SUMIF($B$2:$B$13,B{r},$J$2:$J$13),0)')
        full.cell(r,12,f'=ROUND(K{r}*{[3500,3500,800][i]},0)')
        full.cell(r,13,f'=MIN(E{r},L{r})')
        full.cell(r,14,f'=H{r}*'+(f'IF(E{r}>0,IF(F{r}<7,0.65,1.4),1)' if i==2 else '1'))
        full.cell(r,15,f'=ROUND(M{r}*N{r},2)')
        for col,bom in [(16,[.3,.4,.7][i]),(17,1),(18,[6,7,13][i])]:full.cell(r,col,f'=E{r}*{bom}')
        full.cell(r,19,f'=SUM(O{first}:O{last})')
        full.cell(r,20,f'=ROUND(S{r}*0.2,2)')
        full.cell(r,21,f'=S{r}-T{r}')
        full.cell(r,22,f'=SUM(E{first}:E{last})*2-SUM(M{first}:M{last})*2')
        full.cell(r,23,'=$AB$2');full.cell(r,24,'=$AD$2')

log=book.create_sheet('Control_de_cambios')
log.append(['Hoja','Celda','Antes','Después','Motivo'])
log.append(['FUENTE','SHA256',str(SOURCE.relative_to(ROOT)),hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'Base v16 autorizada; no se encontró v18; original intacto'])
for row in changes:
    # Formula strings in an audit log must remain literal, never execute in this sheet.
    log.append([("'"+v if isinstance(v,str) and v.startswith('=') else v) for v in row])
for ws in [assumptions,oracle,components,full,log]:
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for row in ws:
        for c in row:c.font=Font(name='Arial',size=10)
    for c in ws[1]:c.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=24
assumptions.column_dimensions['C'].width=80;assumptions.column_dimensions['D'].width=80
for col in ['C','D','E']:log.column_dimensions[col].width=65
book.calculation=openpyxl.workbook.properties.CalcProperties(calcId=0,fullCalcOnLoad=True)
DEST.parent.mkdir(exist_ok=True);book.save(DEST)
(ROOT/'reports/excel_v19_build.json').write_text(json.dumps({'source':str(SOURCE.relative_to(ROOT)),'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'output':str(DEST.relative_to(ROOT)),'changes':len(changes),'oracle_cases':len(cases),'oracle_rows':oracle.max_row-1,'recalculated':False},indent=2),encoding='utf8')
print(f'Created {DEST}; {len(changes)} changes, {oracle.max_row-1} oracle rows; RECALC REQUIRED')

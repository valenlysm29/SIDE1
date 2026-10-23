"""Oráculo Excel real. No importa ni ejecuta el motor JavaScript de SIDE.

Uso: python tests/oracle/run_oracle.py [base|sim2|sim3|ruta_escenario.json]
Escenario: {entradas: {Hoja: {A1: valor}}, salidas: {clave: [Hoja, A1]}}.
Los libros originales se copian; errores heredados quedan en el informe JSON.
"""
import argparse
import json
import shutil
import subprocess
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
DIR = ROOT / 'tests/oracle'
TEMP = ROOT / '.codex-work/oracle'
SOURCE = ROOT / '.codex-work/DECISIONES_SIDE_16.xlsx'

def guardar_json(path, obj):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')

def construir():
    TEMP.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, DIR / 'ORACULO_SIDE.xlsx')
    book = openpyxl.load_workbook(DIR / 'ORACULO_SIDE.xlsx')
    # Mantener intactas las diez hojas originales; las extensiones viven aparte.
    sheet = book.create_sheet('Consecuencias_Oracle')
    rows = [
        ['Parámetro','Valor','Regla','Activa','Efecto numérico'],
        ['calidad',12,'CONS-01','=AND(B2<4,B3="SEG-03")','=IF(D2,0.65,1)'],
        ['segmento','SEG-03','CONS-02','=AND(B2>=7,B3="SEG-03")','=IF(D3,1.4,1)'],
        ['eficiencia',17,'CONS-03','=B4<4','=IF(D4,0.7,1)'],
        ['reputacion',8,'CONS-04','=AND(B5+B15<3,B6+1>=2)','=IF(D5,0.8,1)'],
        ['ciclosBajaPrevios',0,'CONS-05','=B7<3','=IF(D6,3,0)'],
        # Matriz_Consecuencias!C7 specifies produced quantity, not planned quantity.
        # The tariff remains an explicit historical extension, not a source formula.
        ['saludFinanciera',-3,'CONS-06','=B11>B10','=IF(D7,0.25*B13*MIN(1,IFERROR((B11-B10)/B10,0)),0)'],
        ['jefatura',True,'CONS-07','=B11>B12','=IF(D8,(B11-B12)*2,0)'],
        ['planeado',520,'CONS-08','=AND(NOT(B8),B4<5)','=IF(D9,0.95,1)'],
        ['capacidad',396],['producido',365],['demanda',350],['costoMO',10000],
        ['produccionTrasConsecuencias','=ROUNDDOWN(MIN(B9,B10)*E4*E9,0)'],
        ['reputacionCambio','=IF(D2,-2,0)+IF(D3,2,0)'],
        ['precioFactor','=E2*E3'],['ciclosBajaFinal','=IF(B5+B15<3,B6+1,0)']]
    for row in rows:
        sheet.append(row)
    for r in range(2,10):
        sheet.cell(r,6, f'=IF(D{r},INDEX(Matriz_Consecuencias!$E$2:$E$9,MATCH(C{r},Matriz_Consecuencias!$A$2:$A$9,0)),"")')
    sheet['F1']='Mensaje literal'
    sheet['A19']='Extensión autorizada: números de horas extra, almacén y tasa son PENDIENTE_DEFINIR del pedido, no fórmulas originales.'
    market = book.create_sheet('Mercado_Oracle')
    market.append(['empresa','segmento','precio','reputacion','calidad','inventario','factorDemanda','mercadoReal','precioIdeal','pesoPrecio','pesoReputacion','pesoCalidad','puntajePrecio','puntajeTotal','porcentaje','asignado','vendido','ingreso','inventarioFinal','precioEfectivo'])
    companies=[['A','SEG-01',40,7,62,1000,1,3500,0,.7,.2,.1],['B','SEG-01',45,8.3,80,1000,1,3500,0,.7,.2,.1],['C','SEG-02',120,10.2,90,1000,1,3500,100,.4,.35,.25],['D','SEG-02',150,9,50,1000,1,3500,100,.4,.35,.25],['E','SEG-03',250,8,100,1000,1,800,250,.1,.5,.4]]
    for r,row in enumerate(companies,2):
        market.append(row)
        maxp=f'_xlfn.MAXIFS($C$2:$C$6,$B$2:$B$6,B{r})'
        minp=f'_xlfn.MINIFS($C$2:$C$6,$B$2:$B$6,B{r})'
        market.cell(r,13,f'=IF(B{r}="SEG-01",IFERROR(100*({maxp}-C{r})/({maxp}-{minp}),100),100-ABS(C{r}-I{r})/I{r}*100)')
        market.cell(r,14,f'=J{r}*M{r}+K{r}*D{r}*10+L{r}*E{r}')
        market.cell(r,15,f'=IFERROR(N{r}/SUMIF($B$2:$B$6,B{r},$N$2:$N$6),0)')
        market.cell(r,16,f'=ROUND(O{r}*H{r},0)')
        market.cell(r,17,f'=MIN(P{r}*G{r},F{r})')
        market.cell(r,18,f'=Q{r}*T{r}')
        market.cell(r,19,f'=F{r}-Q{r}')
        market.cell(r,20,f'=C{r}')
    for sh in [sheet,market]:
        sh.freeze_panes='B2'
        for row in sh:
            for c in row:
                c.font=Font(name='Arial',size=10)
        for c in sh[1]:
            c.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
            c.fill=PatternFill('solid',fgColor='214E65')
        for col in sh.columns:
            sh.column_dimensions[col[0].column_letter].width=22
    # Excel guarda IFS con prefijo; los ArrayFormula se mantienen como tales.
    for sh in book:
        for row in sh:
            for c in row:
                if isinstance(c.value,str) and c.value.startswith('=IFS('):
                    c.value='=_xlfn.'+c.value[1:]
    book.save(DIR/'ORACULO_SIDE.xlsx')
    return book

def escenarios(book):
    prod,cost=book.sheetnames[5:7]
    fixtures=json.loads((ROOT/'docs/logica_excel/fixtures_excel.json').read_text(encoding='utf-8'))
    inputs={'Seleccion_Ciclo0':{f'B{r}':code for r,code in enumerate(fixtures['lotes_motor']['seleccion'],2)},'Lotes_Produccion':{}}
    for r,lot in enumerate(fixtures['lotes_motor']['lotes'],2):
        for col,key in zip('BCDEFG',['segmento','cantidad','cuero','accesorios','hilo','molde']):
            inputs['Lotes_Produccion'][f'{col}{r}']=lot[key]
    outputs={}
    def add(prefix,sh,addresses):
        outputs.update({f'{prefix}.{a}':[sh,a] for a in addresses})
    add('produccion',prod,['I8','I14','I20','K5','L5','M5','N5','L9','L10','M10']+[f'{c}{r}' for r in range(37,40) for c in 'DE']+[f'{c}{r}' for r in range(56,59) for c in 'DEGHI']+[f'{c}{r}' for r in range(63,72) for c in 'DEFG'])
    add('costos',cost,['I8','I14','I20','K5','L5','M5','N5','L9','L10','M10','G65','V95','W95','AB96','AC96','V100','W100']+[f'{c}65' for c in 'CDEF']+[f'{c}{r}' for r in [46,47,48,49] for c in 'DEFHIJLMN']+[f'{c}{r}' for r in range(54,62) for c in 'DEFG'])
    for r in [11,20,29,38]:
        for c in ['U','V','Y','Z','AC','AD']:
            if book[cost][f'{c}{r}'].value is not None:
                add('costos',cost,[f'{c}{r}'])
    for r in [50,57,58,64,66,71]:
        for c in ['U','V','Y','Z','AC','AD']:
            if book[cost][f'{c}{r}'].value is not None and isinstance(book[cost][f'{c}{r}'].value,(int,float)) or (isinstance(book[cost][f'{c}{r}'].value,str) and book[cost][f'{c}{r}'].value.startswith('=')):
                add('costos',cost,[f'{c}{r}'])
    add('lotes','Lotes_Produccion',[f'{c}{r}' for r in range(2,5) for c in 'HIJKL']+['C5','H5'])
    add('motor','Motor_Puntajes',['B6','D6','B9','D9','B10','D10','B11','D11'])
    add('ventas_modelo2',cost,[f'{c}{r}' for start in [118,131,144] for r in range(start,start+4) for c in 'EFGHIJ'])
    add('ventas_modelo1',cost,[f'{c}{r}' for start in [73,87,101] for r in range(start,start+4) for c in 'EFGH'])
    base={'nombre':'base','descripcion':'Fixtures §5: entradas de lotes y selección; hojas originales producción, costos y mercado.','entradas':inputs,'salidas':outputs}
    guardar_json(DIR/'escenario_base.json',base)
    outputs2={f'consecuencias.{book["Consecuencias_Oracle"][f"C{r}"].value}.{key}':['Consecuencias_Oracle',f'{col}{r}'] for r in range(2,10) for key,col in [('activa','D'),('efecto','E'),('mensaje','F')]}
    outputs2.update({f'consecuencias.{key}':['Consecuencias_Oracle',f'B{r}'] for key,r in [('produccion',14),('reputacionCambio',15),('precioFactor',16),('ciclosBajaFinal',17)]})
    sim2={'nombre':'sim2','descripcion':'Escenario numérico con baja calidad Premium, eficiencia baja sin jefe, reputación baja segundo ciclo, salud baja y exceso de plan y producción; CONS-02 en escenario premium separado.','entradas':{'Consecuencias_Oracle':{'B2':3,'B3':'SEG-03','B4':2,'B5':2,'B6':1,'B7':2,'B8':False,'B9':520,'B10':396,'B11':365,'B12':200,'B13':10000}},'salidas':outputs2}
    guardar_json(DIR/'escenario_sim2.json',sim2)
    premium=json.loads(json.dumps(sim2));premium['nombre']='sim2_premium';premium['entradas']['Consecuencias_Oracle']['B2']=12
    guardar_json(DIR/'escenario_sim2_premium.json',premium)
    sim3={'nombre':'sim3','descripcion':'Mercado generalizado cinco empresas; ciclo 1. Para ciclos siguientes, inyectar inventarioInicial y producción en F y factores de demanda en G; pruebas de UI externas.','entradas':{},'salidas':{f'mercado.{r-1}.{key}':['Mercado_Oracle',f'{col}{r}'] for r in range(2,7) for key,col in [('puntajePrecio','M'),('puntajeTotal','N'),('porcentaje','O'),('asignado','P'),('vendido','Q'),('ingreso','R'),('inventarioFinal','S')]}}
    guardar_json(DIR/'escenario_sim3.json',sim3)

def ejecutar(scenario):
    name=scenario.get('nombre','escenario')
    if not name.replace('_','').replace('-','').isalnum():
        raise ValueError('Nombre de escenario no válido')
    plan=TEMP/f'plan_{name}.json'
    guardar_json(plan,scenario)
    output=DIR/f'esperado_{name}.json'
    model=DIR/scenario.get('modelo','ORACULO_SIDE.xlsx')
    subprocess.run(['powershell.exe','-NoProfile','-ExecutionPolicy','Bypass','-File',str(DIR/'recalcular_com.ps1'),'-Plan',str(plan),'-Modelo',str(model),'-LibroSalida',str(TEMP/f'oraculo_{name}.xlsx'),'-JsonSalida',str(output)],cwd=ROOT,check=True,timeout=180)
    return output

def comparacion_base():
    data=json.loads((DIR/'esperado_base.json').read_text(encoding='utf-8-sig'))
    shutil.copy2(TEMP/'oraculo_base.xlsx',DIR/'ORACULO_SIDE_paridad.xlsx')
    book=openpyxl.load_workbook(DIR/'ORACULO_SIDE_paridad.xlsx')
    sheet=book.create_sheet('SIDE vs Excel')
    sheet.append(['Clave','Hoja fuente','Celda fuente','Excel recalculado','SIDE obtenido','Diferencia','Estado'])
    observed_path=DIR/'obtenido_base.json'
    observed=json.loads(observed_path.read_text(encoding='utf-8-sig')) if observed_path.exists() else {}
    observed=observed.get('valores',observed)
    for r,(key,value) in enumerate(data['valores'].items(),2):
        sh,addr=data['celdas'][key]
        sheet.append([key,sh,addr,value,observed.get(key),f'=IF(OR(ISBLANK(D{r}),ISBLANK(E{r})),"",IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),E{r}-D{r},IF(D{r}=E{r},0,1)))',f'=IF(ISBLANK(E{r}),"PENDIENTE SIDE",IF(ABS(F{r})<=1E-9*MAX(1,ABS(IFERROR(D{r}*1,0))),"OK","DIFERENCIA"))'])
    sheet.freeze_panes='D2'
    for col,width in [('A',42),('B',28),('C',15),('D',24),('E',24),('F',20),('G',24)]:sheet.column_dimensions[col].width=width
    book.save(DIR/'ORACULO_SIDE_paridad.xlsx')
    # Guardar cálculos de comparación mediante el mismo Excel real.
    ejecutar_comparacion={'nombre':'comparacion','entradas':{},'salidas':{}}
    plan=TEMP/'plan_comparacion.json';guardar_json(plan,ejecutar_comparacion)
    subprocess.run(['powershell.exe','-NoProfile','-ExecutionPolicy','Bypass','-File',str(DIR/'recalcular_com.ps1'),'-Plan',str(plan),'-Modelo',str(DIR/'ORACULO_SIDE_paridad.xlsx'),'-LibroSalida',str(TEMP/'comparacion.xlsx'),'-JsonSalida',str(TEMP/'comparacion.json')],cwd=ROOT,check=True,timeout=180)
    shutil.copy2(TEMP/'comparacion.xlsx',DIR/'ORACULO_SIDE_paridad.xlsx')

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('escenario',nargs='?',default='base');parser.add_argument('--construir',action='store_true');parser.add_argument('--comparacion',action='store_true');args=parser.parse_args()
    if args.construir or not (DIR/'ORACULO_SIDE.xlsx').exists():escenarios(construir())
    path=Path(args.escenario)
    if not path.exists():path=DIR/f'escenario_{args.escenario}.json'
    result=ejecutar(json.loads(path.read_text(encoding='utf-8-sig')))
    if args.comparacion:comparacion_base()
    print(result.relative_to(ROOT))

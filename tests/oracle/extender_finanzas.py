"""Añade fórmulas contables independientes a las hojas pipeline recalculables."""
import json
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter as L
DIR=Path(__file__).resolve().parent

def extender(escenario,config):
    name=escenario['id'].lower().replace('-','');cfg={**config,**escenario.get('overrides',{})}
    path=DIR/f'ORACULO_{name}.xlsx';book=openpyxl.load_workbook(path);sh=book['Pipeline_'+name]
    heads={c.value:c.column for c in sh[1] if c.value is not None};plan_path=DIR/f'escenario_{name}_pipeline.json';plan=json.loads(plan_path.read_text(encoding='utf8'))
    def col(k):
        if k not in heads:heads[k]=len(heads)+1;sh.cell(1,heads[k],k)
        return L(heads[k])
    def a(k,r):return f'{col(k)}{r}'
    def formula(k,r,expr):sh[a(k,r)]='='+expr
    def value(k,r,v):sh[a(k,r)]=v
    previous={};assets_prev={}
    for r in range(2,sh.max_row+1):
        id=sh[a('empresa',r)].value;cycle=sh[a('ciclo',r)].value;prev=previous.get(id);R=lambda k:a(k,r);P=lambda k:a(k,prev) if prev else '0'
        phase=escenario['ciclos'][cycle-1]['decisionesPorEmpresa'];d=phase.get(id)
        if d is None:
            d=next(p['decisionesPorEmpresa'][id] for p in reversed(escenario['ciclos'][:cycle-1]) if id in p['decisionesPorEmpresa'])
        machinery=d.get('maquinaria',{});owned=assets_prev.get(id,{})
        retained={**owned};purchase=[];acquisition=[]
        rates=cfg['TASAS_ACTIVOS']['NO_MYPE' if d.get('regimen')=='NO_MYPE' else 'MYPE']
        for code,count in machinery.items():
            retained[code]=max(owned.get(code,0),count);delta=max(0,retained[code]-owned.get(code,0));eq=cfg['MAQUINARIA'][code]
            installation=0 if eq.get('molde') else rates['instalacion'];training=rates['capacitacionMolde'] if eq.get('molde') else rates['capacitacion']
            purchase.append(f'{delta}*{eq["costo"]}*(1+{installation}+{training})');acquisition.append(f'{delta}*{eq["costo"]}')
        assets_prev[id]=retained
        capital=d.get('capitalInicial',escenario.get('capitalInicial',1000000));value('capitalInicial',r,capital)
        formula('compraActivos',r,'+'.join(purchase) or '0');formula('adquisicionNeta',r,'+'.join(acquisition) or '0')
        factor=(1-cfg['DESCUENTO_CANTIDAD'])*(1-cfg['DESCUENTO_PROVEEDOR_UNICO'] if 'INS-09A' in d['seleccion'] else 1)*(1-cfg['DESCUENTO_CONTADO'] if 'INS-10A' in d['seleccion'] else 1)
        if d.get('analistaCompras'):
            analyst=cfg['ANALISTA_BASE'];factor*=1-max(analyst['descuentoMinimoPct'],analyst['descuentoInicialPct']-(cycle-1)*analyst['reduccionPorCicloPct'])/100
        formula('valorCompradoNeto',r,'('+('+'.join(f'{R("comprado_"+c)}*{price}' for c,price in cfg['PRECIO_MP'].items()))+f')*{factor}')
        formula('valorInventarioNeto',r,'('+('+'.join(f'{R("stockMP_"+c)}*{price}' for c,price in cfg['PRECIO_MP'].items()))+f')*{factor}')
        formula('cajaInicial',r,P('cajaFinal') if prev else R('capitalInicial'))
        formula('tasaNuevoPrestamo',r,f'(1+{R("tasaTEA")}/100)^(1/{cfg["MESES_ANIO"]})-1')
        formula('interesMensual',r,f'{P("deudaPrincipal")}*{P("tasaMensual")}+{R("montoPrestamo")}*{R("tasaNuevoPrestamo")}')
        formula('tasaMensual',r,f'IFERROR({R("interesMensual")}/({P("deudaPrincipal")}+{R("montoPrestamo")}),0)')
        formula('intereses',r,f'{R("interesMensual")}*{cfg["CICLO_EQUIVALE_A_MESES"]}' if cfg['PRESTAMO_COBRAR_INTERES_CICLO'] else '0')
        formula('pagosOperacion',r,f'{R("costosTotales")}-{R("depreciacion")}-{R("costoMP")}+{R("valorCompradoNeto")}+{R("intereses")}')
        formula('cajaFinal',r,f'{R("cajaInicial")}+{R("ingresos")}+{R("montoPrestamo")}-{R("pagosOperacion")}-{R("compraActivos")}')
        formula('depreciacionAcumulada',r,f'{P("depreciacionAcumulada")}+{R("depreciacion")}')
        formula('activosCosto',r,f'{P("activosCosto")}+{R("adquisicionNeta")}')
        formula('activosFijos',r,f'{R("activosCosto")}-{R("depreciacionAcumulada")}')
        mapped={
            'estadoResultados.ingresos':'ingresos','estadoResultados.ventasNetas':'ingresos','estadoResultados.costos':f'{R("costosTotales")}+{R("intereses")}',
            'estadoResultados.utilidad':f'{R("utilidad")}-{R("intereses")}','estadoResultados.depreciacion':'depreciacion','estadoResultados.intereses':'intereses','estadoResultados.costosProduccion':'costosTotales',
            'balanceCaja.cajaInicial':'cajaInicial','balanceCaja.entradas':f'{R("ingresos")}+{R("montoPrestamo")}','balanceCaja.salidas':f'{R("pagosOperacion")}+{R("compraActivos")}','balanceCaja.cajaFinal':'cajaFinal',
            'flujoCaja.cajaInicial':'cajaInicial','flujoCaja.cobrosVentas':'ingresos','flujoCaja.pagosOperacion':'pagosOperacion','flujoCaja.compraActivos':'compraActivos','flujoCaja.financiamiento':'montoPrestamo',
            'flujoCaja.operacion':f'{R("ingresos")}-{R("pagosOperacion")}','flujoCaja.inversion':f'-{R("compraActivos")}','flujoCaja.flujoNeto':f'{R("cajaFinal")}-{R("cajaInicial")}','flujoCaja.cajaFinal':'cajaFinal',
            'balanceGeneral.efectivo':'cajaFinal','balanceGeneral.activosFijos':'activosFijos','balanceGeneral.inventarioMP':'valorInventarioNeto',
            'balanceGeneral.activos':f'{R("cajaFinal")}+{R("activosFijos")}+{R("valorInventarioNeto")}',
            'balanceGeneral.deuda':'deudaPrincipal','balanceGeneral.patrimonio':f'{R("cajaFinal")}+{R("activosFijos")}+{R("valorInventarioNeto")}-{R("deudaPrincipal")}',
            'balanceGeneral.pasivoPatrimonio':f'{R("cajaFinal")}+{R("activosFijos")}+{R("valorInventarioNeto")}',
            'financiamiento.montoPrestamo':'montoPrestamo','financiamiento.tasaTEA':'tasaTEA','financiamiento.deudaPrincipal':'deudaPrincipal','financiamiento.tasaAdicionalProximoPrestamo':'penalizacionPendiente','financiamiento.tasaMensual':'tasaMensual','financiamiento.intereses':'intereses'}
        for field,expr in mapped.items():
            key='financial.'+field;formula(key,r,R(expr) if expr in heads else expr)
            plan['salidas'][f'ciclo.{cycle}.{id}.{key}']=[sh.title,R(key)]
        previous[id]=r
    book.save(path);plan_path.write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf8');print(name,len(plan['salidas']))

if __name__=='__main__':
    data=json.loads((DIR/'entradas_simulaciones.json').read_text(encoding='utf8'))
    for name in ['sim1','sim2','sim3']:
        extender(json.loads((DIR/f'entradas_{name}.json').read_text(encoding='utf8')),data['config'])
    for path in ['entradas_banca_recuperacion.json','entradas_banca_mezcla.json']:
        bank=json.loads((DIR/path).read_text(encoding='utf8'));extender(bank['escenario'],data['config'])

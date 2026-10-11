"""Extiende el pipeline con lotes múltiples, FIFO e identidades entre ciclos."""
import json
import copy
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter as L
from construir_pipeline import construir,DIR

def construir_multilote(original,config):
    cfg={**config,**original.get('overrides',{})}
    surrogate=copy.deepcopy(original);surrogate['permitir_identidades']=True
    for phase in surrogate['ciclos']:
        for d in phase['decisionesPorEmpresa'].values():
            total=sum(x['cantidad'] for x in d['lotes']);d['lotes']=[{**d['lotes'][0],'cantidad':total}]
    name,plan=construir(surrogate,config)
    book=openpyxl.load_workbook(DIR/f'ORACULO_{name}.xlsx')
    company=book['Pipeline_'+name];lots=book.create_sheet('Lotes_Pipeline_'+name);market=book.create_sheet('Mercados_Pipeline_'+name)
    hc={c.value:c.column for c in company[1] if c.value is not None};hl={};hm={}
    def ensure(sh,heads,k):
        if k not in heads:heads[k]=len(heads)+1;sh.cell(1,heads[k],k)
        return L(heads[k])
    def A(sh,heads,key,row):return f'{ensure(sh,heads,key)}{row}'
    def put(sh,heads,key,row,value):sh[A(sh,heads,key,row)]=value
    def CA(key,row):return A(company,hc,key,row)
    def CR(key,row):return f"'{company.title}'!{CA(key,row)}"
    def LA(key,row):return A(lots,hl,key,row)
    def LR(key,row):return f"'{lots.title}'!{LA(key,row)}"
    def MA(key,row):return A(market,hm,key,row)
    def MR(key,row):return f"'{market.title}'!{MA(key,row)}"
    def CF(key,row,formula):put(company,hc,key,row,'='+formula)
    def LF(key,row,formula):put(lots,hl,key,row,'='+formula)
    def MF(key,row,formula):put(market,hm,key,row,'='+formula)
    def look(code,column):return f'IFERROR(INDEX(\'Catalogo_Numerico_{name}\'!${column}$2:${column}$69,MATCH({code},\'Catalogo_Numerico_{name}\'!$A$2:$A$69,0)),0)'
    codes=list(cfg['PRECIO_MP']);families=cfg['INSUMOS_POR_FAMILIA'];catalog_rows={};next_lot=2;next_market=2
    identities={id:{} for id in original['empresas']};previous={};company_groups=[]
    for cycle,phase in enumerate(original['ciclos'],1):
        for id,d in phase['decisionesPorEmpresa'].items():
            cr=next(r for r in range(2,company.max_row+1) if company[CA('ciclo',r)].value==cycle and company[CA('empresa',r)].value==id)
            actual={tuple(x.get(k,'') for k in ['segmento','molde','cuero','accesorios','hilo']):x for x in d['lotes']}
            for key,lot in actual.items():
                if key not in identities[id]:identities[id][key]=lot
            lot_rows=[];previous_lot=None;active_segments=[]
            for key,stored in identities[id].items():
                lot=actual.get(key);lr=next_lot;next_lot+=1;before=previous.get((id,key))
                for key2,val in [('ciclo',cycle),('empresa',id),('segmento',stored['segmento']),('molde',stored['molde']),('cuero',stored.get('cuero','')),('accesorios',stored.get('accesorios','')),('hilo',stored.get('hilo','')),('cantidad',lot['cantidad'] if lot else 0),('precio',lot.get('precio',0) if lot else stored.get('precio',0))]:put(lots,hl,key2,lr,val)
                LF('calidad',lr,'+'.join(look(LA(k,lr),'H') for k in ['cuero','accesorios','hilo','molde']))
                cons=cfg['CONSUMO_MOLDE'][stored['molde']]
                for code in codes:
                    expr=LR('saldoPre_'+code,previous_lot) if previous_lot else CR('disponible_'+code,cr)
                    LF('disponible_'+code,lr,expr)
                limits=[]
                for family,famcodes in families.items():
                    eligible=[stored[family]] if stored.get(family) else famcodes
                    limits.append('('+('+'.join(LA('disponible_'+c,lr) for c in eligible))+f')/{cons[family]}')
                LF('capacidadAsignada',lr,f'IFERROR(ROUNDDOWN({LA("cantidad",lr)}*{CR("produccionCapacidad",cr)}/{CR("planeado",cr)},0),0)')
                LF('producidoPre',lr,f'ROUNDDOWN(MIN({LA("cantidad",lr)},{LA("capacidadAsignada",lr)},{",".join(limits)}),0)')
                LF('producido',lr,f'ROUNDDOWN({LA("producidoPre",lr)}*{CR("factorConsecuencias",cr)},0)')
                for family,famcodes in families.items():
                    used=[]
                    for code in famcodes:
                        expr='0'
                        if not stored.get(family) or stored[family]==code:
                            need=f'{LA("producidoPre",lr)}*{cons[family]}'+('-(%s)'%'+'.join(used) if used else '')
                            expr=f'MIN({LA("disponible_"+code,lr)},MAX(0,{need}))'
                        LF('usadoPre_'+code,lr,expr);used.append(LA('usadoPre_'+code,lr))
                        LF('saldoPre_'+code,lr,f'{LA("disponible_"+code,lr)}-{LA("usadoPre_"+code,lr)}')
                        LF('usado_'+code,lr,f'IFERROR({LA("usadoPre_"+code,lr)}*{LA("producido",lr)}/{LA("producidoPre",lr)},0)')
                LF('precioFactor',lr,f'IF(AND({LA("segmento",lr)}="SEG-03",{LA("cantidad",lr)}>0),IF({LA("calidad",lr)}<4,0.65,IF({LA("calidad",lr)}>=7,1.4,1)),1)')
                # Precio efectivo de inventario conservado; decisión actual de misma identidad lo actualiza.
                LF('precioEfectivo',lr,f'{LA("precio",lr)}*{LA("precioFactor",lr)}' if lot else LR('precioEfectivo',before))
                LF('stockAntes',lr,f'{LR("inventarioFinal",before) if before else "0"}+{LA("producido",lr)}')
                previous[(id,key)]=lr;previous_lot=lr;lot_rows.append(lr)
                if stored['segmento'] not in active_segments:active_segments.append(stored['segmento'])
            current_rows=[row for row in lot_rows if lots[LA('cantidad',row)].value>0]
            CF('calidad',cr,'IFERROR(('+('+'.join(f'{LR("cantidad",row)}*{LR("calidad",row)}' for row in current_rows))+f')/{CA("planeado",cr)},0)')
            bad='OR('+','.join(f'AND({LR("segmento",row)}="SEG-03",{LR("cantidad",row)}>0,{LR("calidad",row)}<4)' for row in current_rows)+')'
            good='OR('+','.join(f'AND({LR("segmento",row)}="SEG-03",{LR("cantidad",row)}>0,{LR("calidad",row)}>=7)' for row in current_rows)+')'
            CF('reputacionCambio',cr,f'IF({bad},-2,0)+IF({good},2,0)');CF('CONS-01',cr,bad);CF('CONS-02',cr,good)
            CF('producido',cr,'+'.join(LR('producido',row) for row in lot_rows))
            for family,famcodes in families.items():CF('consumo'+family.capitalize(),cr,'+'.join(CA('usado_'+code,cr) for code in famcodes))
            for code in codes:
                CF('usado_'+code,cr,'+'.join(LR('usado_'+code,row) for row in lot_rows))
                CF('stockMP_'+code,cr,f'{CA("disponible_"+code,cr)}-{CA("usado_"+code,cr)}')
            proportions=[]
            for family,famcodes in families.items():
                denominator=famcodes[1:] if family=='hilo' and cfg.get('REPUTACION_HILO_DENOMINADOR_LITERAL') else famcodes
                proportions.append(f'IFERROR({CA("usado_"+famcodes[-1],cr)}/({"+".join(CA("usado_"+c,cr) for c in denominator)}),0)')
            premium='+'.join(LR('producido',row) for row in lot_rows if lots[LA('molde',row)].value=='INF-10') or '0'
            CF('reputacionOperativa',cr,f'IFERROR(({premium})/{CA("producido",cr)},0)*10+IFERROR({CA("n3",cr)}/{CA("trabajadores",cr)},0)*10+({"+".join(proportions)})/3*10+{CA("publicidad",cr)}/10000*10')
            market_rows=[]
            for segment in cfg['MERCADOS']:
                relevant=[row for row in lot_rows if lots[LA('segmento',row)].value==segment]
                if not relevant:continue
                current=[row for row in current_rows if row in relevant];mr=next_market;next_market+=1
                m=cfg['MERCADOS'][segment]
                for k,v in [('ciclo',cycle),('empresa',id),('segmento',segment),('grupo',f'{cycle}|{segment}')]:put(market,hm,k,mr,v)
                stock='+'.join(LR('stockAntes',row) for row in relevant)
                MF('activo',mr,'TRUE()' if current else f'({stock})>0')
                weighted='+'.join(f'{LR("precio",row)}*{LR("cantidad" if current else "stockAntes",row)}' for row in (current or relevant))
                denom='+'.join(LR('cantidad' if current else 'stockAntes',row) for row in (current or relevant))
                MF('precio',mr,f'IFERROR(({weighted})/({denom}),0)')
                MF('reputacion',mr,CR('reputacionMercado',cr));MF('calidad',mr,f'ROUND(MIN(100,{CR("calidad",cr)}*10),0)')
                market_rows.append({'row':mr,'segment':segment,'m':m,'current':bool(current)})
                sold_prior=[]
                for row in relevant:
                    LF('vendido',row,f'MIN({LA("stockAntes",row)},MAX(0,{MR("demandaAjustada",mr)}'+('-(%s)'%'+'.join(sold_prior) if sold_prior else '')+'))')
                    sold_prior.append(LR('vendido',row))
                    LF('ingreso',row,f'{LA("vendido",row)}*{LA("precioEfectivo",row)}')
                    LF('inventarioFinal',row,f'{LA("stockAntes",row)}-{LA("vendido",row)}')
            first=next(x['row'] for x in market_rows if x['current'])
            CF('asignado',cr,MR('asignado',first));CF('demandaAjustada',cr,MR('demandaAjustada',first))
            CF('vendido',cr,'+'.join(LR('vendido',row) for row in lot_rows));CF('ingresos',cr,'+'.join(LR('ingreso',row) for row in lot_rows))
            CF('inventarioProducto',cr,'+'.join(LR('inventarioFinal',row) for row in lot_rows))
            company_groups.append({'row':cr,'markets':market_rows,'selection':d['seleccion'],'cycle':cycle,'id':id,'current_lots':current_rows})
    last=market.max_row
    for group in company_groups:
        cr=group['row']
        for item in group['markets']:
            r=item['row'];R=lambda key:MA(key,r);m=item['m']
            gr=f'${ensure(market,hm,"grupo")}$2:${ensure(market,hm,"grupo")}${last}';pr=f'${ensure(market,hm,"precio")}$2:${ensure(market,hm,"precio")}${last}';active=f'${ensure(market,hm,"activo")}$2:${ensure(market,hm,"activo")}${last}'
            maxp=f'_xlfn.MAXIFS({pr},{gr},{R("grupo")},{active},TRUE())';minp=f'_xlfn.MINIFS({pr},{gr},{R("grupo")},{active},TRUE())'
            pp=f'IFERROR(100*({maxp}-{R("precio")})/({maxp}-{minp}),100)' if item['segment']=='SEG-01' else f'100-ABS({R("precio")}-{m["precioIdeal"]})/{m["precioIdeal"]}*100'
            MF('puntajePrecio',r,pp);MF('puntajeTotal',r,f'IF({R("activo")},{m["pesos"][0]}*{R("puntajePrecio")}+{m["pesos"][1]}*{R("reputacion")}*10+{m["pesos"][2]}*{R("calidad")},0)')
            tr=f'${ensure(market,hm,"puntajeTotal")}$2:${ensure(market,hm,"puntajeTotal")}${last}'
            if cfg.get('CALIDAD_MERCADO_MODO')=='excel_literal':
                rr=f'${ensure(market,hm,"reputacion")}$2:${ensure(market,hm,"reputacion")}${last}'
                MF('calidad',r,f'IFERROR(ROUND({R("reputacion")}/SUMIFS({rr},{gr},{R("grupo")},{active},TRUE())*100,0),0)')
            MF('reputacionEscalada',r,f'{R("reputacion")}*10')
            MF('porcentaje',r,f'IFERROR({R("puntajeTotal")}/SUMIF({gr},{R("grupo")},{tr}),0)')
            MF('asignado',r,f'IFERROR(ROUND({R("puntajeTotal")}/SUMIF({gr},{R("grupo")},{tr})*{m["unidades"]}*{m["penetracion"]},0),0)')
            MF('demandaAjustada',r,f'MAX(0,ROUNDDOWN({R("asignado")}*{CR("factorDemanda",cr)}*{1+cfg["BONO_DEMANDA_MARKETING"] if "INV-02B" in group["selection"] else 1},0))')
            if name=='sim1':
                for field,colkey in [('puntajePrecio','puntajePrecio'),('puntajeTotal','puntajeTotal'),('porcentaje','porcentaje'),('asignado','asignado'),('demandaAjustada','demandaAjustada'),('calidadMercado','calidad'),('reputacionEscalada','reputacionEscalada')]:
                    plan['salidas'][f'ciclo.{group["cycle"]}.{group["id"]}.mercado.{item["segment"]}.{field}']=[market.title,R(colkey)]
        if name=='sim1':
            for index,lr in enumerate(group['current_lots']):
                for field,colkey in [('calidad','calidad'),('producido','producido'),('capacidadAsignada','capacidadAsignada'),('precioEfectivo','precioEfectivo')]:
                    plan['salidas'][f'ciclo.{group["cycle"]}.{group["id"]}.lote.{index}.{field}']=[lots.title,LA(colkey,lr)]
    for sh in [lots,market]:
        sh.freeze_panes='D2'
        for col in range(1,sh.max_column+1):sh.column_dimensions[L(col)].width=20
    book.save(DIR/f'ORACULO_{name}.xlsx')
    (DIR/f'escenario_{name}_pipeline.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
    print(name,len(plan['salidas']),'salidas,',len(company_groups),'empresa-ciclos,',lots.max_row-1,'filas lote')
    return name,plan

if __name__=='__main__':
    import sys
    selected=sys.argv[1] if len(sys.argv)>1 else 'sim2'
    data=json.loads((DIR/'entradas_simulaciones.json').read_text(encoding='utf-8-sig'))
    original=json.loads((DIR/f'entradas_{selected}.json').read_text(encoding='utf-8-sig'))
    construir_multilote(original,data['config'])

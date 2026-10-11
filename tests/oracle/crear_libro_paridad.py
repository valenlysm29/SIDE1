"""Añade comparaciones revisables a una copia de las hojas fuente originales."""
import json
import shutil
import subprocess
from pathlib import Path
import openpyxl
from openpyxl.styles import Font,PatternFill

ROOT=Path(__file__).resolve().parents[2];DIR=ROOT/'tests/oracle';TEMP=ROOT/'.codex-work/oracle'
def read(path):return json.loads(path.read_text(encoding='utf-8-sig'))

def main():
    book=openpyxl.load_workbook(TEMP/'oraculo_base.xlsx')
    summary=book.create_sheet('SIDE vs Excel')
    summary.append(['Escenario','Celdas comparadas','OK','Diferencias','Pendientes','Origen esperado'])
    cases=[('base','Base - 351'),('sim1_pipeline','SIM1 - Pipeline'),('sim2_pipeline','SIM2 - Pipeline'),('sim3_pipeline','SIM3 - Pipeline'),('sim2','CONS baja'),('sim2_premium','CONS premium')]
    for summary_row,(name,title) in enumerate(cases,2):
        ep=DIR/f'esperado_{name}.json'
        if not ep.exists():continue
        expected=read(ep);op=DIR/f'obtenido_{name}.json';observed=read(op) if op.exists() else {};observed=observed.get('valores',observed)
        sheet=book.create_sheet(title)
        sheet.append(['Clave','Hoja de origen','Celda','Excel recalculado','SIDE obtenido','Diferencia','Estado','SIDE recibido'])
        for r,(key,value) in enumerate(expected['valores'].items(),2):
            source,addr=expected['celdas'][key]
            sheet.append([key,source,addr,value,observed.get(key),f'=IF(NOT(H{r}),"",IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),E{r}-D{r},IF(D{r}=E{r},0,1)))',f'=IF(NOT(H{r}),"PENDIENTE SIDE",IF(IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),IF(MOD(D{r},1)=0,D{r}=E{r},ABS(D{r}-E{r})<=1E-9*MAX(1,ABS(D{r}),ABS(E{r}))),D{r}=E{r}),"OK","DIFERENCIA"))',key in observed])
        last=sheet.max_row
        summary.append([name,last-1,f'=COUNTIF(\'{title}\'!G2:G{last},"OK")',f'=COUNTIF(\'{title}\'!G2:G{last},"DIFERENCIA")',f'=COUNTIF(\'{title}\'!G2:G{last},"PENDIENTE SIDE")',expected['metadatos']['motor']])
        sheet.freeze_panes='D2';sheet.auto_filter.ref=f'A1:H{last}'
        for col,width in [('A',42),('B',28),('C',14),('D',24),('E',24),('F',20),('G',22)]:sheet.column_dimensions[col].width=width
    summary.append(['NOTA','Las cuatro divisiones por cero originales de CANTIDAD se conservan y quedan fuera de las comparaciones.'])
    summary.append(['Motor','Microsoft Excel COM 16.0; LibreOffice no instalado.'])
    summary.append(['Independencia','Las hojas de pipeline se construyen exclusivamente desde decisiones y parámetros; los obtenidos SIDE solo llenan las columnas de comparación.'])
    for col,width in [('A',28),('B',28),('C',20),('D',20),('E',20),('F',32)]:summary.column_dimensions[col].width=width
    for sheet in list(book)[book.index(summary):]:
        for row in sheet:
            for c in row:c.font=Font(name='Arial',size=10)
        for c in sheet[1]:c.font=Font(name='Arial',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='214E65')
    book.active=book.index(summary)
    target=DIR/'ORACULO_SIDE_paridad.xlsx';book.save(target)
    plan=TEMP/'plan_libro_paridad.json';plan.write_text(json.dumps({'nombre':'libro_paridad','entradas':{},'salidas':{}}),encoding='utf-8')
    subprocess.run(['powershell.exe','-NoProfile','-ExecutionPolicy','Bypass','-File',str(DIR/'recalcular_com.ps1'),'-Plan',str(plan),'-Modelo',str(target),'-LibroSalida',str(TEMP/'libro_paridad.xlsx'),'-JsonSalida',str(TEMP/'libro_paridad.json')],cwd=ROOT,check=True,timeout=180)
    shutil.copy2(TEMP/'libro_paridad.xlsx',target)
    check=openpyxl.load_workbook(target,data_only=True)
    print([(check['SIDE vs Excel'].cell(r,1).value,[check['SIDE vs Excel'].cell(r,c).value for c in range(2,6)]) for r in range(2,2+sum((DIR/f'esperado_{n}.json').exists() for n,_ in cases))])

if __name__=='__main__':main()

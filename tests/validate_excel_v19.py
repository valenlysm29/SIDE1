"""Independent structural/error audit; run from SIDE1, never saves the source."""
import json
import re
import sys
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
source = root / (sys.argv[1] if len(sys.argv) > 1 else 'docs/DECISIONES_SIDE_19.xlsx')
formula = load_workbook(source, data_only=False)
cached = load_workbook(source, data_only=True)
errors, truncated, dirty, dates, missing_cache = [], [], [], [], []
for sheet in formula:
    for row in sheet:
        for cell in row:
            value = cell.value
            text = getattr(value, 'text', value)
            cache = cached[sheet.title][cell.coordinate]
            address = f'{sheet.title}!{cell.coordinate}'
            if cache.data_type == 'e' or (isinstance(cache.value, str) and re.match(r'^#(?:DIV/0!|REF!|N/A|VALUE!|NAME\?|NUM!|NULL!)$', cache.value)):
                errors.append({'cell': address, 'value': cache.value})
            if cell.data_type == 'f' and cache.value is None:
                missing_cache.append(address)
            if isinstance(text, str) and cell.data_type == 'f' and re.search(r'\$[A-Z]+\$2:\$[A-Z]+\$57', text):
                truncated.append(address)
            if sheet.title == 'Catalogo_Decisiones' and cell.row > 1 and cell.column in range(7, 12) and isinstance(value, str) and cell.data_type != 'f':
                if re.search(r'S/|^=\+', value):
                    dirty.append({'cell': address, 'value': value})
            if sheet.title == 'CANTIDAD DE PRODUCCIÓN' and cell.coordinate in ('K56','K57','K58') and isinstance(value, datetime):
                dates.append(address)
catalog_count = sum(bool(formula['Catalogo_Decisiones'].cell(row, 1).value) for row in range(2, formula['Catalogo_Decisiones'].max_row+1))
result = {'source': str(source.relative_to(root)), 'catalog_options': catalog_count, 'formula_errors': errors, 'truncated_lookups': truncated, 'dirty_numeric_cells': dirty, 'corrupted_dates': dates, 'missing_formula_caches': missing_cache,
          'has_change_control': 'Control_de_cambios' in formula.sheetnames, 'has_assumptions': 'Supuestos' in formula.sheetnames}
out = root/'tests/output'/('validate-'+source.stem+'.json')
out.parent.mkdir(parents=True, exist_ok=True)
out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({key: len(value) if isinstance(value, list) else value for key, value in result.items()}, ensure_ascii=False))
if '_19' in source.stem:
    assert catalog_count == 68
    assert not errors and not truncated and not dirty and not dates
    assert result['has_change_control'] and result['has_assumptions']

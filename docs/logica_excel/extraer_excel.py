"""Extrae el original sin modificarlo. Ejecutar desde la raíz SIDE1."""
import hashlib
import json
import re
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / '.codex-work/DECISIONES_SIDE_16.xlsx'
OUT = ROOT / 'docs/logica_excel'
formulas = openpyxl.load_workbook(SOURCE, data_only=False)
values = openpyxl.load_workbook(SOURCE, data_only=True)

def write(name, data):
    (OUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')

def formula_text(cell):
    return getattr(cell.value, 'text', cell.value)

dump = {}
for sheet in formulas:
    cells = []
    for row in sheet:
        for c in row:
            if c.value is None:
                continue
            item = {'celda': c.coordinate, 'contenido': formula_text(c),
                    'valor_cache': values[sheet.title][c.coordinate].value,
                    'tipo': c.data_type, 'formato': c.number_format}
            if hasattr(c.value, 'ref'):
                item['rango_matricial'] = c.value.ref
            cells.append(item)
    dump[sheet.title] = cells
    safe = re.sub(r'[^\w-]', '_', sheet.title)
    (OUT / ('volcado_' + safe + '.txt')).write_text(
        'Hoja: ' + sheet.title + '\nCelda\tContenido/fórmula literal\tValor en caché\tTipo\tRango matricial\n' +
        '\n'.join('\t'.join(json.dumps(x.get(k), ensure_ascii=False, default=str) for k in
                            ['celda', 'contenido', 'valor_cache', 'tipo', 'rango_matricial']) for x in cells) + '\n', encoding='utf-8')
write('celdas_excel.json', dump)

catalog = []
keys = ['codigo', 'bloque', 'grupo', 'opcion', 'permanencia', 'descripcion', 'costo',
        'calidad', 'eficiencia', 'reputacion', 'financiero', 'notas']
for r in range(2, values['Catalogo_Decisiones'].max_row + 1):
    row = [values['Catalogo_Decisiones'].cell(r, c).value for c in range(1, 13)]
    if not row[0]:
        continue
    item = dict(zip(keys, row))
    item['costo_original'] = item['costo']
    if isinstance(item['costo'], str):
        item['costo'] = float(item['costo'].replace('S/', '').replace(',', '').strip())
    item['fila_excel'] = r
    catalog.append(item)
write('catalogo_excel.json', catalog)

def cell(s, addr):
    return values[s][addr].value

def block(s, bounds):
    return {c.coordinate: c.value for row in values[s][bounds] for c in row}

prod = formulas.sheetnames[5]
cost = formulas.sheetnames[6]
RECALCULADO = ROOT / '.codex-work/DECISIONES_SIDE_16_recalculado.xlsx'
if RECALCULADO.exists():
    values = openpyxl.load_workbook(RECALCULADO, data_only=True)
estado_verificacion = 'recalculado_excel_com_16' if RECALCULADO.exists() else 'cache_original'
selection = ['INF-03','INF-06','INF-09','RH-02','RH-05','RH-08','RH-10A','COM-01C','FIN-01B','INS-09B','INS-10A','INS-11B','INV-01B','INV-02B']
lots = [
    {'segmento':'SEG-03','cantidad':100,'cuero':'INS-03','accesorios':'INS-05','hilo':'INS-08','molde':'INF-10'},
    {'segmento':'SEG-02','cantidad':300,'cuero':'INS-02','accesorios':'INS-04','hilo':'INS-06','molde':'INF-08'},
    {'segmento':'SEG-03','cantidad':100,'cuero':'INS-01','accesorios':'INS-04','hilo':'INS-06','molde':'INF-08'}]
by_code = {x['codigo']:x for x in catalog}
for lot in lots:
    lot['costoUnitario'] = sum(by_code[lot[k]]['costo'] for k in ['cuero','accesorios','hilo'])
    lot['calidad'] = sum(by_code[lot[k]]['calidad'] for k in ['cuero','accesorios','hilo','molde'])
    lot['umbral'] = {'SEG-01':0,'SEG-02':4,'SEG-03':7}[lot['segmento']]
    lot['resultado'] = 'Coherente' if lot['calidad'] >= lot['umbral'] else 'No alcanza el segmento'

fixtures = {
    'metadatos': {'fuente':'DECISIONES SIDE (16).xlsx', 'sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
                  'hojas':formulas.sheetnames, 'nota':'Caches originales no equivalen a recálculo; lote_motor verificado por catálogo y fórmulas, pendiente confirmar con LibreOffice.'},
    'lotes_motor': {'seleccion':selection,'lotes':lots,'cantidadTotal':500,'costoTotalInsumos':28000,
                    'calidadEmpresa':6.2,'nivelCalidad':'Medio',
                    **{key:sum(by_code[c][col] for c in selection) for key,col in [('eficiencia','eficiencia'),('reputacion','reputacion'),('saludFinanciera','financiero')]},
                    'verificacion':'derivado_independiente_de_catalogo_y_formulas_no_recalculado'},
    'produccion': {'hoja':prod, 'dias':24, 'capacidadAreas':[cell(prod,a) for a in ['I8','I14','I20']],
                   'capacidadDiaria':cell(prod,'K5'),'capacidadCiclo':cell(prod,'L5'),
                   'factorEficiencia':cell(prod,'M5'),'produccionCapacidad':cell(prod,'N5'),
                   'trabajadores':cell(prod,'L9'),'trabajadoresN3':cell(prod,'L10'),'proporcionN3':cell(prod,'M10'),
                   'planeado':[cell(prod,f'D{r}') for r in range(37,40)],
                   'reparto':[cell(prod,f'E{r}') for r in range(37,40)],
                   'real':[cell(prod,f'E{r}') for r in range(56,59)],
                   'rangos':{r:block(prod,r) for r in ['D31:L33','D49:L51','D56:I58','D63:G71']},
                   'verificacion':'cache_original'},
    'costos': {'hoja':cost,'dias':30,'capacidadAreas':[cell(cost,a) for a in ['I8','I14','I20']],
               'capacidadDiaria':cell(cost,'K5'),'capacidadCiclo':cell(cost,'L5'),'factorEficiencia':cell(cost,'M5'),
               'produccionCapacidad':cell(cost,'N5'),'trabajadores':cell(cost,'L9'),'trabajadoresN3':cell(cost,'L10'),
               'proporcionN3':cell(cost,'M10'),'produccionReal':[cell(cost,f'E{r}') for r in range(46,49)],
               'reputacionOperativa':cell(cost,'G65'),'factoresReputacion':[cell(cost,f'{c}65') for c in 'CDEF'],
               'valorCompradoLista':cell(cost,'V95'),'valorCompradoNeto':cell(cost,'W95'),
               'valorConsumidoLista':cell(cost,'AB96'),'valorConsumidoNeto':cell(cost,'AC96'),
               'valorInventarioLista':cell(cost,'V100'),'valorInventarioNeto':cell(cost,'W100'),
               'rangos':{r:block(cost,r) for r in ['T4:AD39','T44:AD71','D31:K33','D46:N49','D54:H61','T76:AD100']},
               'verificacion':'cache_original'},
    'mercado_modelo2':{},'mercado_modelo1':{},
    'consecuencias':[dict(zip(['id','puntaje','condicion','efecto','mensaje'],[values['Matriz_Consecuencias'].cell(r,c).value for c in range(1,6)])) for r in range(2,10)]
}
for segment,start,market in [('economico',118,3500),('estandar',131,3500),('premium',144,800)]:
    rows = list(range(start,start+4))
    fixtures['mercado_modelo2'][segment] = {'mercadoReal':market,
        **{key:[cell(cost,f'{col}{r}') for r in rows] for key,col in [('precios','C'),('reputaciones','D'),('calidades','E'),('puntajesPrecio','F'),('puntajesTotal','H'),('unidades','J')]},'verificacion':'cache_original'}
for segment,start,market in [('economico',73,3500),('estandar',87,3500),('premium',101,800)]:
    fixtures['mercado_modelo1'][segment] = {'mercadoReal':market,
        **{key:[cell(cost,f'{col}{r}') for r in range(start,start+4)] for key,col in [('precios','D'),('reputaciones','C'),('puntajesPrecio','E'),('puntajesTotal','F'),('unidades','H')]},'verificacion':'cache_original'}
write('fixtures_excel.json', fixtures)
if RECALCULADO.exists():
    fixtures['metadatos']['nota'] = 'Producción/costos/ventas leídos de copia recalculada por Microsoft Excel COM 16.0. Lotes/Motor necesitan escenario con entradas.'
    for key in ['produccion','costos']:
        fixtures[key]['verificacion'] = estado_verificacion
    for key in ['mercado_modelo1','mercado_modelo2']:
        for entry in fixtures[key].values():
            entry['verificacion'] = estado_verificacion
    write('fixtures_excel.json', fixtures)
oracle_base = ROOT / 'tests/oracle/esperado_base.json'
if oracle_base.exists():
    recalculadas = json.loads(oracle_base.read_text(encoding='utf-8-sig'))['valores']
    fixtures['lotes_motor']['verificacion'] = 'recalculado_excel_com_16'
    fixtures['lotes_motor']['celdasRecalculadas'] = {k:v for k,v in recalculadas.items() if k.startswith(('lotes.','motor.'))}
    write('fixtures_excel.json', fixtures)
write('metadatos_excel.json',{'sha256':fixtures['metadatos']['sha256'],
    'hojas':[{'nombre':s.title,'filas':s.max_row,'columnas':s.max_column,'celdasNoVacias':len(dump[s.title]),
              'formulas':sum(1 for c in dump[s.title] if c['tipo']=='f'),
              'matriciales':sum(1 for c in dump[s.title] if 'rango_matricial' in c)} for s in formulas],
    'catalogoFilas':len(catalog),
    'errores_cache_original':[{'hoja':s,'celda':x['celda'],'valor':x['valor_cache']} for s,cells in dump.items() for x in cells if isinstance(x['valor_cache'],str) and x['valor_cache'].startswith(('#DIV/0!','#REF!','#VALUE!','#NAME?','#N/A','#NUM!','#NULL!'))]})
print(json.dumps({'hojas':len(dump),'catalogo':len(catalog),'lotes_motor':fixtures['lotes_motor'],
                  'produccion':fixtures['produccion']['real'],'reputacionOperativa':fixtures['costos']['reputacionOperativa']},ensure_ascii=True))
